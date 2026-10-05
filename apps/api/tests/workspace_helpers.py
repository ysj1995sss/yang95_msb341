"""Shared setup for the workspace API tests (spec 009). Synthetic data only."""

import io
from datetime import datetime, timedelta, timezone

from docx import Document
from jose import jwt

SECRET = "w" * 40


def resume_docx() -> bytes:
    doc = Document()
    for line in ("Riley Park", "riley@example.com | 555-0100 | Denver, CO", "", "EXPERIENCE", "Data Analyst",
                 "Acme Analytics | Denver, CO | 2021 - Present"):
        doc.add_paragraph(line)
    for bullet in ("Built SQL dashboards used by 40 managers", "Cut weekly reporting time by 30%"):
        doc.add_paragraph(f"• {bullet}")
    for line in ("", "EDUCATION", "BS Economics, State University, 2019", "", "SKILLS", "SQL, Tableau, Python"):
        doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def token(sub="google-123", name="Riley Park", secret=SECRET, minutes=5, audience="job-copilot-workspace"):
    claims = {"sub": sub, "name": name, "aud": audience,
              "exp": datetime.now(timezone.utc) + timedelta(minutes=minutes)}
    return jwt.encode(claims, secret, algorithm="HS256")
