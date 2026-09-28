"""Steps 16-20 (Task 8 groundwork): regenerating an artifact from reviewed
ResumeChange dispositions, without ever re-invoking the LLM tailorer."""

import io

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from resume_tailorer.analyzers.gap_analyzer import GapReport
from resume_tailorer.artifacts.models import ChangeCategory, ChangeDisposition, ResumeChange, ValidationStatus
from resume_tailorer.artifacts.regeneration import (
    apply_dispositions_to_text,
    final_text_for_change,
    regenerate_docx_artifact,
    validate_manual_text,
)
from resume_tailorer.models import CareerTruthProfile, WorkExperience


def _change(
    change_id="c1", original_text="Original bullet", proposed_text="Proposed bullet",
    disposition=ChangeDisposition.PENDING, source_index=None, manual_text=None,
) -> ResumeChange:
    return ResumeChange(
        change_id=change_id, section="work_experience", source_index=source_index,
        original_text=original_text, proposed_text=proposed_text,
        category=ChangeCategory.REPHRASED, reason="test", job_requirement="",
        evidence_source="career_profile", evidence_text=original_text,
        validation_status=ValidationStatus.PASS, disposition=disposition, manual_text=manual_text,
    )


class TestFinalTextForChange:
    def test_pending_keeps_the_proposed_text(self):
        change = _change(disposition=ChangeDisposition.PENDING)
        assert final_text_for_change(change) == "Proposed bullet"

    def test_accepted_keeps_the_proposed_text(self):
        change = _change(disposition=ChangeDisposition.ACCEPTED)
        assert final_text_for_change(change) == "Proposed bullet"

    def test_rejected_reverts_to_original(self):
        change = _change(disposition=ChangeDisposition.REJECTED)
        assert final_text_for_change(change) == "Original bullet"

    def test_restored_reverts_to_original(self):
        change = _change(disposition=ChangeDisposition.RESTORED)
        assert final_text_for_change(change) == "Original bullet"

    def test_manually_edited_uses_manual_text_not_proposed_text(self):
        """manual_text is kept SEPARATE from proposed_text (which stays
        the untouched AI proposal) -- see ResumeChange.manual_text's
        docstring. An earlier version of this code overwrote proposed_text
        directly, which broke apply_dispositions_to_text's find-and-replace
        (it could no longer locate the AI's original wording in the
        baseline text) and silently dropped every manual edit on
        regeneration."""
        change = _change(
            proposed_text="AI's original proposal", manual_text="User's own edit",
            disposition=ChangeDisposition.MANUALLY_EDITED,
        )
        assert final_text_for_change(change) == "User's own edit"

    def test_manually_edited_without_manual_text_falls_back_to_proposed_text(self):
        change = _change(proposed_text="User's own edit", disposition=ChangeDisposition.MANUALLY_EDITED)
        assert final_text_for_change(change) == "User's own edit"


class TestApplyDispositionsToText:
    def test_rejected_change_restores_the_original_bullet(self):
        baseline = "- Built REST APIs\n- Deployed cloud services on AWS infrastructure"
        change = _change(
            original_text="Deployed services on AWS",
            proposed_text="Deployed cloud services on AWS infrastructure",
            disposition=ChangeDisposition.REJECTED,
        )
        result = apply_dispositions_to_text(baseline, [change])
        assert "Deployed cloud services on AWS infrastructure" not in result
        assert "Deployed services on AWS" in result

    def test_pending_change_is_left_as_proposed(self):
        baseline = "- Deployed cloud services on AWS infrastructure"
        change = _change(
            original_text="Deployed services on AWS",
            proposed_text="Deployed cloud services on AWS infrastructure",
            disposition=ChangeDisposition.PENDING,
        )
        result = apply_dispositions_to_text(baseline, [change])
        assert result == baseline

    def test_rebuilding_from_baseline_is_idempotent(self):
        """Flipping a disposition back and forth must never compound --
        every regeneration rebuilds from the ORIGINAL baseline, never from
        a previous regeneration's output."""
        baseline = "- Deployed cloud services on AWS infrastructure"
        change = _change(
            original_text="Deployed services on AWS",
            proposed_text="Deployed cloud services on AWS infrastructure",
            disposition=ChangeDisposition.REJECTED,
        )
        first = apply_dispositions_to_text(baseline, [change])
        second = apply_dispositions_to_text(first, [change])
        assert first == second == "- Deployed services on AWS"

    def test_manually_edited_change_is_substituted_into_the_baseline(self):
        """Regression test: apply_dispositions_to_text must find the AI's
        ORIGINAL proposal (proposed_text, untouched) in the baseline text
        and replace it with manual_text -- not compare proposed_text
        against itself, which is what silently dropped every manual edit
        before ResumeChange.manual_text was split out as its own field."""
        baseline = "- Deployed cloud services on AWS infrastructure"
        change = _change(
            original_text="Deployed services on AWS",
            proposed_text="Deployed cloud services on AWS infrastructure",
            manual_text="Deployed cloud services on AWS infrastructure (reviewed by a human)",
            disposition=ChangeDisposition.MANUALLY_EDITED,
        )
        result = apply_dispositions_to_text(baseline, [change])
        assert result == "- Deployed cloud services on AWS infrastructure (reviewed by a human)"


class TestValidateManualText:
    def _profile(self):
        return CareerTruthProfile(
            contact_info={}, education=[],
            work_experience=[
                WorkExperience(
                    employer="Acme", title="Engineer", dates="2020-2023",
                    responsibilities=[], accomplishments=["Deployed services on AWS"],
                )
            ],
            skills=["Python"], tools=[], certifications=[], accomplishments=[],
        )

    def test_safe_rephrase_passes(self):
        issues = validate_manual_text(
            "Deployed services on AWS", "Deployed services using AWS", self._profile()
        )
        assert issues == []

    def test_fabricated_claim_is_flagged(self):
        issues = validate_manual_text(
            "Deployed services on AWS", "Led a team of 50 engineers deploying services on AWS",
            self._profile(),
        )
        assert issues


def _add_bullet_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph(text, style="List Paragraph")
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    return paragraph


def _sample_docx_bytes(tmp_path) -> bytes:
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    doc.add_paragraph("Marketing Manager")
    doc.add_paragraph("Acme Corp | Remote\t2022-2024")
    _add_bullet_paragraph(doc, "Did a thing")
    path = tmp_path / "sample.docx"
    doc.save(path)
    return path.read_bytes()


class TestRegenerateDocxArtifact:
    def test_rejected_change_reverts_the_docx_paragraph(self, tmp_path):
        original_bytes = _sample_docx_bytes(tmp_path)
        doc = Document(io.BytesIO(original_bytes))
        target_index = next(i for i, p in enumerate(doc.paragraphs) if p.text == "Did a thing")

        change = _change(
            original_text="Did a thing", proposed_text="Did a rewritten thing",
            disposition=ChangeDisposition.REJECTED, source_index=target_index,
        )
        profile = CareerTruthProfile(
            contact_info={}, education=[], work_experience=[], skills=[], tools=[],
            certifications=[], accomplishments=[],
        )

        result = regenerate_docx_artifact(
            original_docx_bytes=original_bytes, changes=[change], profile=profile,
            gap_report=GapReport(items=[], summary="none"), convert_to_pdf=False,
        )

        spliced = Document(io.BytesIO(result.docx_bytes))
        assert any(p.text == "Did a thing" for p in spliced.paragraphs)
        assert not any(p.text == "Did a rewritten thing" for p in spliced.paragraphs)
