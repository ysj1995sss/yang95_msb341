"""
Integration tests for the resume tailoring pipeline (Task 13).

SCOPE RULING — read before extending this file
================================================
The plan's Task 13 spec calls for validating the MVP's Global Constraint:
"80% keyword alignment on 8/10 real job postings." That goal requires:
  1. Real job postings (not accessible from this environment), and
  2. Live LLM calls via ResumeTailorer.tailor() and
     ResumeTailoringOptimizer.optimize() (no provider API key is configured
     in this environment, and these tests must run in CI/sandboxes without
     network access or API cost).

Because of that, `ResumeTailorer.tailor()` and `ResumeTailoringOptimizer.optimize()`
are NOT exercised against a real model here. The 80% alignment goal is NOT
validated by this test file and must be validated manually, or in a follow-up
task that has live API access, by running the full pipeline (see Test 3 below,
with mocking removed) against a set of real job postings and confirming
optimizer.final_score >= 0.80 on at least 8 of 10 postings.

What IS covered here (no API needed):
  - Test 1: resume parsing produces a valid CareerTruthProfile.
  - Test 2: for each of 10 representative job descriptions, the
    analyze -> benchmark -> gap_analysis chain runs without exceptions and
    produces well-formed output (this exercises Tasks 4-7, including the
    keyword-alignment scoring used by the optimizer in Task 6).
  - Test 3: a single mocked end-to-end run (parse -> analyze -> benchmark ->
    gap -> tailor[mocked] -> optimize[mocked] -> PDF generate -> PDF validate
    -> report) verifying that data flows correctly between every stage of
    the pipeline, injecting a MagicMock LLMClient. This test
    checks integration/wiring, not alignment quality.
"""

from pathlib import Path
from unittest.mock import MagicMock

import pytest

from resume_tailorer.models import CareerTruthProfile, WorkExperience
from resume_tailorer.parsers.resume_parser import ResumeParser
from resume_tailorer.analyzers.job_analyzer import JobAnalyzer, JobAnalysis
from resume_tailorer.analyzers.resume_benchmarker import ResumeBenchmarker, ResumeBenchmark
from resume_tailorer.analyzers.gap_analyzer import GapAnalyzer, GapReport
from resume_tailorer.llm.client import LLMClient
from resume_tailorer.tailorer.resume_tailorer import ResumeTailorer
from resume_tailorer.tailorer.optimizer import ResumeTailoringOptimizer, OptimizationResult
from resume_tailorer.pdf.generator import PDFGenerator
from resume_tailorer.pdf.validator import PDFValidator, ValidationResult
from resume_tailorer.report_generator import ReportGenerator
from resume_tailorer.docx_export.pipeline import run_docx_tailoring_pipeline
from resume_tailorer.tailorer.docx_bullet_tailorer import BulletEdit, BulletTailoringResult
from resume_tailorer.artifacts.models import FidelityMode
from resume_tailorer.artifacts.report import build_final_report

from tests.fixtures.real_job_descriptions import JOB_DESCRIPTIONS


SAMPLE_RESUME_PATH = Path(__file__).parent / "fixtures" / "sample_resume.pdf"


def _mock_llm(return_text: str) -> MagicMock:
    """Build a MagicMock LLMClient with complete() returning return_text."""
    llm = MagicMock(spec=LLMClient)
    llm.complete.return_value = return_text
    return llm


@pytest.fixture(scope="module")
def parsed_profile():
    """Parse the sample resume fixture once for the whole module."""
    parser = ResumeParser()
    return parser.parse(str(SAMPLE_RESUME_PATH))


class TestResumeParsingIntegration:
    """Test 1: resume parsing produces a valid, well-formed CareerTruthProfile."""

    def test_parses_sample_resume_into_valid_profile(self, parsed_profile):
        assert isinstance(parsed_profile, CareerTruthProfile)
        assert isinstance(parsed_profile.contact_info, dict)
        assert parsed_profile.name  # non-empty via the .name property
        assert isinstance(parsed_profile.skills, list)
        assert isinstance(parsed_profile.tools, list)
        assert isinstance(parsed_profile.certifications, list)
        assert isinstance(parsed_profile.work_experience, list)
        assert isinstance(parsed_profile.education, list)


class TestJobDescriptionPipelineIntegration:
    """
    Test 2: for each of the 10 representative job descriptions, run
    analyze -> benchmark -> gap_analysis and verify no exceptions and
    reasonable output shapes. No Claude API calls are involved in this
    chain (JobAnalyzer, ResumeBenchmarker, and GapAnalyzer are all
    regex/heuristic-based per Tasks 4, 5, and 7).
    """

    @pytest.mark.parametrize("job_description", JOB_DESCRIPTIONS)
    def test_full_analysis_chain_runs_without_exceptions(self, parsed_profile, job_description):
        job_analyzer = JobAnalyzer()
        benchmarker = ResumeBenchmarker()
        gap_analyzer = GapAnalyzer()

        # Stage 1: analyze the job description.
        job_analysis = job_analyzer.analyze(job_description)
        assert isinstance(job_analysis, JobAnalysis)
        assert isinstance(job_analysis.required_qualifications, list)
        assert isinstance(job_analysis.preferred_qualifications, list)
        assert isinstance(job_analysis.responsibilities, list)
        assert isinstance(job_analysis.skills_required, list)
        assert isinstance(job_analysis.tools_required, list)
        assert isinstance(job_analysis.weighted_keywords, list)

        # Stage 2: benchmark the original (untailored) resume against the job.
        benchmark = benchmarker.benchmark(parsed_profile, job_analysis)
        assert isinstance(benchmark, ResumeBenchmark)
        assert 0.0 <= benchmark.original_match_score <= 1.0
        assert isinstance(benchmark.keywords_matched, list)
        assert isinstance(benchmark.keywords_missing, list)
        # Every required skill/tool keyword must land in exactly one bucket.
        total_keywords = len(job_analysis.skills_required) + len(job_analysis.tools_required)
        assert len(benchmark.keywords_matched) + len(benchmark.keywords_missing) == total_keywords

        # Stage 3: gap analysis between the profile and the job.
        gap_report = gap_analyzer.analyze(parsed_profile, job_analysis, benchmark)
        assert isinstance(gap_report, GapReport)
        assert isinstance(gap_report.items, list)
        assert isinstance(gap_report.summary, str)
        assert gap_report.summary  # non-empty
        # Every gap item must reference one of the A-E categories with a reason.
        for item in gap_report.items:
            assert item.requirement
            assert item.category is not None
            assert isinstance(item.reason, str) and item.reason

    def test_ten_job_descriptions_are_present(self):
        """Sanity check that the fixture file actually provides 10 postings."""
        assert len(JOB_DESCRIPTIONS) == 10
        assert all(isinstance(jd, str) and jd.strip() for jd in JOB_DESCRIPTIONS)


class TestMockedEndToEndPipeline:
    """
    Test 3: one mocked full-pipeline run covering every stage, verifying
    that data flows correctly end to end without errors. This does NOT
    validate alignment quality (that requires the live API — see the
    module docstring) — it only validates that each component's output is
    shaped correctly for the next component's input.
    """

    def test_full_pipeline_parse_through_report(self, tmp_path):
        # Mocked LLM response used for both tailor() and optimize()'s
        # internal refine step.
        tailored_content = (
            "Jane Doe\n"
            "SUMMARY\n"
            "Experienced Software Engineer with Python, Docker, and PostgreSQL background.\n"
            "WORK EXPERIENCE\n"
            "Software Engineer | Acme Corp | 2018-2023\n"
            "- Built APIs and maintained services\n"
            "- Reduced latency by 30%\n"
        )
        llm = _mock_llm(tailored_content)

        # Stage 1: parse.
        parser = ResumeParser()
        profile = parser.parse(str(SAMPLE_RESUME_PATH))
        assert isinstance(profile, CareerTruthProfile)

        # Stage 2: analyze one representative job description.
        job_analyzer = JobAnalyzer()
        job_analysis = job_analyzer.analyze(JOB_DESCRIPTIONS[0])
        assert isinstance(job_analysis, JobAnalysis)

        # Stage 3: benchmark original resume.
        benchmarker = ResumeBenchmarker()
        benchmark = benchmarker.benchmark(profile, job_analysis)
        assert isinstance(benchmark, ResumeBenchmark)

        # Stage 4: gap analysis.
        gap_analyzer = GapAnalyzer()
        gap_report = gap_analyzer.analyze(profile, job_analysis, benchmark)
        assert isinstance(gap_report, GapReport)

        # Stage 5: tailor (LLM call mocked).
        tailorer = ResumeTailorer(llm=llm)
        tailored_text = tailorer.tailor(profile, job_analysis, gap_report)
        assert isinstance(tailored_text, str) and tailored_text
        assert llm.complete.called

        # Stage 6: optimize (LLM call mocked, capped to 1 iteration so
        # the test stays fast and deterministic).
        optimizer = ResumeTailoringOptimizer(max_iterations=1, llm=llm)
        optimization_result = optimizer.optimize(
            profile, job_analysis, tailored_text, gap_report
        )
        assert isinstance(optimization_result, OptimizationResult)
        assert isinstance(optimization_result.tailored_resume, str)
        assert isinstance(optimization_result.final_score, float)

        # Stage 7: generate PDF from the optimized resume text.
        pdf_generator = PDFGenerator()
        output_path = str(tmp_path / "tailored_resume.pdf")
        generated_path = pdf_generator.generate(
            optimization_result.tailored_resume,
            profile.name,
            output_path=output_path,
        )
        assert Path(generated_path).exists()

        # Stage 8: validate the generated PDF (round-trip hard gate).
        validator = PDFValidator()
        validation_result = validator.validate(generated_path)
        assert isinstance(validation_result, ValidationResult)
        assert validation_result.passed is True
        assert validation_result.issues == []

        # Stage 9: build the final report from every prior stage's output.
        report = ReportGenerator().generate_report(
            profile,
            job_analysis,
            benchmark,
            gap_report,
            optimization_result,
            validation_result,
        )
        assert isinstance(report, dict)
        expected_keys = {
            "candidate_name",
            "job_title",
            "candidate_fit_score",
            "candidate_fit_score_available",
            "original_match_score",
            "tailored_match_score",
            "score_improvement",
            "optimization_iterations",
            "optimization_ceiling_reached",
            "qualifications_summary",
            "gaps_addressed_count",
            "total_gaps_count",
            "pdf_validation_passed",
            "pdf_validation_details",
            "recommendations",
        }
        assert expected_keys.issubset(report.keys())
        assert report["candidate_name"] == profile.name
        assert report["pdf_validation_passed"] is True


def _sample_docx_bytes(tmp_path) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    doc.add_paragraph("Marketing Manager")
    doc.add_paragraph("Acme Corp | Remote\t2022-2024")
    paragraph = doc.add_paragraph("Did a thing", style="List Paragraph")
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    path = tmp_path / "sample.docx"
    doc.save(path)
    return path.read_bytes()


class _StubBulletTailorer:
    """Rewrites only the one bullet -- exercises the real splice/validate
    path without an LLM call."""

    def tailor_bullets(self, bullets, profile, job_analysis, gap_report, max_repair_attempts=1, priority_focus=None):
        edits = []
        for b in bullets:
            if b.text == "Did a thing":
                edits.append(BulletEdit(b.paragraph_index, b.text, "Did a thing using SQL", changed=True))
            else:
                edits.append(BulletEdit(b.paragraph_index, b.text, b.text, changed=False))
        return BulletTailoringResult(edits=edits, warnings=[])


class TestArtifactPipelineIntegration:
    """Task 8 (spec 002/decisions 014): app.py and the API router now
    dispatch a DOCX upload to run_docx_tailoring_pipeline -- and its
    output feeds build_final_report -- instead of the legacy freeform
    reconstruction path Test 3 above covers. This is a mocked end-to-end
    run of THAT path (no LLM call), verifying the artifact/report models
    stay wired together correctly."""

    def test_docx_upload_produces_a_validated_report(self, tmp_path):
        profile = CareerTruthProfile(
            contact_info={"name": "Jane Doe"},
            education=[],
            work_experience=[
                WorkExperience(
                    employer="Acme Corp", title="Marketing Manager", dates="2022-2024",
                    responsibilities=[], accomplishments=[],
                )
            ],
            skills=[], tools=[], certifications=[], accomplishments=[],
        )
        job_analysis = JobAnalyzer().analyze("Marketing role requiring SQL.")
        benchmark = ResumeBenchmarker().benchmark(profile, job_analysis)
        gap_report = GapAnalyzer().analyze(profile, job_analysis, benchmark)

        original_bytes = _sample_docx_bytes(tmp_path)
        docx_result = run_docx_tailoring_pipeline(
            original_bytes, profile, job_analysis, gap_report,
            bullet_tailorer=_StubBulletTailorer(), convert_to_pdf=False,
        )

        assert docx_result.docx_bytes
        assert docx_result.changes
        assert docx_result.validation.status is not None

        report = build_final_report(
            candidate_fit=None,
            fit_breakdown={},
            original_alignment=benchmark.original_match_score,
            tailored_alignment=benchmark.original_match_score,
            gap_report=gap_report,
            validation=docx_result.validation,
            fidelity_mode=FidelityMode.PRESERVED,
        )

        assert report.fidelity_mode is FidelityMode.PRESERVED
        assert report.candidate_fit is None
        assert report.validation is docx_result.validation
