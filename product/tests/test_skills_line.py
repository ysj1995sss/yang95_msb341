"""Spec 022: the skills line in a Word resume, where change levels 2-3 may add skills.
Synthetic documents only."""

from docx import Document

from resume_tailorer.docx_export.skills_line import append_skills, find_skills_paragraph_index


def _doc(*lines):
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    return doc


def test_the_line_under_a_skills_heading_is_found():
    doc = _doc("Sam Example", "EXPERIENCE", "Data Analyst | Acme | 2021 - Present", "EDUCATION",
               "BS Economics, State University, 2019", "SKILLS", "SQL, Excel, Tableau, Python")
    assert find_skills_paragraph_index(doc.paragraphs) == 6


def test_an_inline_skills_label_is_found():
    doc = _doc("Sam Example", "Technical Skills: SQL, Python, Tableau", "EXPERIENCE")
    assert find_skills_paragraph_index(doc.paragraphs) == 1


def test_no_list_means_no_skills_line():
    doc = _doc("Sam Example", "SKILLS", "I am good at many things and enjoy learning")
    assert find_skills_paragraph_index(doc.paragraphs) is None
    assert find_skills_paragraph_index(_doc("Sam Example", "EXPERIENCE").paragraphs) is None


def test_two_skills_headings_are_ambiguous_and_left_alone():
    doc = _doc("SKILLS", "SQL, Excel", "TECHNICAL SKILLS", "Python, R")
    assert find_skills_paragraph_index(doc.paragraphs) is None


def test_append_keeps_the_lines_own_separator_and_label():
    assert append_skills("SQL, Excel, Tableau", ["PostgreSQL", "Looker"]) == "SQL, Excel, Tableau, PostgreSQL, Looker"
    assert append_skills("SQL | Excel", ["Looker"]) == "SQL | Excel | Looker"
    assert append_skills("Technical Skills: SQL; Python", ["pandas"]) == "Technical Skills: SQL; Python; pandas"
    assert append_skills("SQL, Excel.", ["Looker"]) == "SQL, Excel, Looker."
