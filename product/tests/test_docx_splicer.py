from docx import Document

from resume_tailorer.tailorer.docx_bullet_tailorer import BulletEdit
from resume_tailorer.docx_export.splicer import (
    inline_formatting_findings,
    splice_bullets_into_docx,
)


def test_multi_run_paragraph_collapses_to_run_zero_formatting():
    doc = Document()
    p = doc.add_paragraph()
    run1 = p.add_run("Grew")
    run1.font.name = "Times New Roman"
    run1.bold = True
    run2 = p.add_run(" repeat bookings by 20%")
    run2.font.name = "Times New Roman"
    run2.bold = True

    edits = [BulletEdit(paragraph_index=0, original_text=p.text, new_text="Grew repeat bookings by 25%", changed=True)]
    spliced = splice_bullets_into_docx(doc, edits)

    paragraph = spliced.paragraphs[0]
    assert paragraph.text == "Grew repeat bookings by 25%"
    assert len(paragraph.runs) == 1
    assert paragraph.runs[0].font.name == "Times New Roman"
    assert paragraph.runs[0].bold is True


def test_unchanged_paragraph_is_never_touched():
    doc = Document()
    p = doc.add_paragraph()
    run1 = p.add_run("Untouched")
    run2 = p.add_run(" bullet text")
    original_run_count = len(p.runs)

    edits = [BulletEdit(paragraph_index=0, original_text=p.text, new_text=p.text, changed=False)]
    spliced = splice_bullets_into_docx(doc, edits)

    assert len(spliced.paragraphs[0].runs) == original_run_count
    assert spliced.paragraphs[0].runs[0].text == "Untouched"
    assert spliced.paragraphs[0].runs[1].text == " bullet text"


def test_paragraph_with_no_runs_gets_a_new_run_added():
    doc = Document()
    p = doc.add_paragraph()  # no runs at all
    edits = [BulletEdit(paragraph_index=0, original_text="", new_text="New content", changed=True)]
    spliced = splice_bullets_into_docx(doc, edits)
    assert spliced.paragraphs[0].text == "New content"


def test_only_targeted_paragraph_is_modified_others_untouched():
    doc = Document()
    doc.add_paragraph("Untouched heading")
    p1 = doc.add_paragraph("Original bullet one")
    doc.add_paragraph("Untouched bullet two")

    edits = [BulletEdit(paragraph_index=1, original_text=p1.text, new_text="Rewritten bullet one", changed=True)]
    spliced = splice_bullets_into_docx(doc, edits)

    assert spliced.paragraphs[0].text == "Untouched heading"
    assert spliced.paragraphs[1].text == "Rewritten bullet one"
    assert spliced.paragraphs[2].text == "Untouched bullet two"


def test_mixed_inline_formatting_on_changed_bullet_warns():
    doc = Document()
    paragraph = doc.add_paragraph(style="List Paragraph")
    paragraph.add_run("Led ")
    paragraph.add_run("critical").bold = True
    paragraph.add_run(" launch")

    findings = inline_formatting_findings(paragraph, changed=True)

    assert [finding.code for finding in findings] == ["INLINE_FORMATTING_SIMPLIFIED"]


def test_mixed_inline_formatting_on_unchanged_bullet_does_not_warn():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Led ")
    paragraph.add_run("critical").bold = True

    assert inline_formatting_findings(paragraph, changed=False) == []
