"""Typed bullet glyphs and one-line job headers (decision 027). Synthetic data only."""

import pytest
from docx import Document

from resume_tailorer.docx_export.splicer import inline_formatting_findings, splice_bullets_into_docx
from resume_tailorer.parsers.bullets import strip_typed_bullet, typed_bullet_prefix
from resume_tailorer.parsers.docx_structure import extract_docx_structure, is_bullet_paragraph
from resume_tailorer.parsers.resume_parser import ResumeParser, split_one_line_header
from resume_tailorer.tailorer.docx_bullet_tailorer import BulletEdit


@pytest.mark.parametrize("text, prefix", [
    ("• Led a team", "• "),
    ("▪\tBuilt dashboards", "▪\t"),
    ("- Wrote SQL", "- "),
    ("* Shipped", "* "),
    ("-5% churn", ""),
    ("office manager", ""),
    ("•", ""),
])
def test_typed_bullet_prefix(text, prefix):
    assert typed_bullet_prefix(text) == prefix


def _typed_resume():
    doc = Document()
    doc.add_paragraph("Riley Park")
    doc.add_paragraph("riley@example.com")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("Data Analyst")
    doc.add_paragraph("Acme Corp | Denver, CO\t2021 - Present")
    glyph_run_bullet = doc.add_paragraph()
    marker = glyph_run_bullet.add_run("•\t")
    marker.font.name = "Symbol"
    body = glyph_run_bullet.add_run("Built SQL dashboards for 40 managers")
    body.font.name = "Calibri"
    doc.add_paragraph("• Cleaned data in Excel")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("BS Economics, State U, 2018")
    return doc


def test_typed_bullets_are_splice_targets_without_their_glyph():
    doc = _typed_resume()
    structure = extract_docx_structure(doc)
    assert len(structure.jobs) == 1
    targets = [b for b in structure.splice_targets(doc.paragraphs) if b.section == "work_experience"]
    assert [b.text for b in targets] == ["Built SQL dashboards for 40 managers", "Cleaned data in Excel"]


def test_splicing_a_typed_bullet_keeps_the_marker_and_its_font():
    doc = _typed_resume()
    edits = [
        BulletEdit(5, "Built SQL dashboards for 40 managers", "Built SQL and Tableau dashboards for 40 managers", True),
        BulletEdit(6, "Cleaned data in Excel", "Cleaned and validated data in Excel", True),
    ]
    assert inline_formatting_findings(doc.paragraphs[5], True) == []
    splice_bullets_into_docx(doc, edits)
    first, second = doc.paragraphs[5], doc.paragraphs[6]
    assert first.text == "•\tBuilt SQL and Tableau dashboards for 40 managers"
    assert first.runs[0].font.name == "Symbol" and first.runs[1].font.name == "Calibri"
    assert second.text == "• Cleaned and validated data in Excel"
    assert is_bullet_paragraph(second)


def test_docx_text_extraction_reads_typed_bullets_once(tmp_path):
    path = tmp_path / "typed.docx"
    _typed_resume().save(path)
    profile = ResumeParser().parse(str(path))
    job = profile.work_experience[0]
    bullets = job.responsibilities + job.accomplishments
    assert "Cleaned data in Excel" in bullets
    assert all(not b.startswith(("•", "-")) for b in bullets)


@pytest.mark.parametrize("line, expected", [
    ("Data Analyst | Acme Corp | Denver, CO | Jan 2021 - Present", ("Data Analyst", "Acme Corp", "Denver, CO", "Jan 2021 - Present")),
    ("Acme Corp | Marketing Manager | May 2019 - Dec 2020", ("Marketing Manager", "Acme Corp", "", "May 2019 - Dec 2020")),
    ("Senior Engineer at Globex, Remote 2018-2020", ("Senior Engineer", "Globex", "Remote", "2018-2020")),
    ("Acme Corp 2019 - 2020", None),
])
def test_split_one_line_header(line, expected):
    assert split_one_line_header(line) == expected


def test_one_line_headers_and_other_glyphs_in_pdf_text():
    text = (
        "Riley Park\nriley@example.com\n\nEXPERIENCE\n"
        "Data Analyst | Acme Corp | Denver, CO | Jan 2021 - Present\n"
        "▪ Built SQL dashboards used by 40 managers\n"
        "▪ Cut report time by 30%\n"
        "Junior Analyst | Globex | 2019 - 2020\n"
        "* Cleaned data in Excel\n"
        "\nEDUCATION\nBS Economics, State U, 2018\n"
    )
    jobs = ResumeParser()._parse_text(text).work_experience
    assert [(j.title, j.employer, j.dates) for j in jobs] == [
        ("Data Analyst", "Acme Corp", "Jan 2021 - Present"),
        ("Junior Analyst", "Globex", "2019 - 2020"),
    ]
    assert jobs[0].accomplishments == ["Cut report time by 30%"]
    assert jobs[1].responsibilities == ["Cleaned data in Excel"]


def test_two_line_headers_still_work():
    text = (
        "Riley Park\n\nEXPERIENCE\nData Analyst\nAcme Corp | Denver, CO\t2021 - Present\n"
        "• Built dashboards\n\nEDUCATION\nBS Economics, 2018\n"
    )
    job = ResumeParser()._parse_text(text).work_experience[0]
    assert (job.title, job.employer, job.location) == ("Data Analyst", "Acme Corp", "Denver, CO")


def test_strip_typed_bullet_leaves_plain_text_alone():
    assert strip_typed_bullet("Plain sentence") == "Plain sentence"


def test_a_wrapped_bullet_before_a_one_line_header_stays_with_its_job():
    text = (
        "Riley Park\n\nEXPERIENCE\n"
        "Data Analyst | Acme Corp | Jan 2021 - Present\n"
        "• Built SQL dashboards used by 40 regional managers across the\n"
        "sales organization\n"
        "Junior Analyst | Globex | 2019 - 2020\n"
        "• Cleaned data in Excel\n"
        "\nEDUCATION\nBS Economics, State U, 2018\n"
    )
    jobs = ResumeParser()._parse_text(text).work_experience
    assert [j.title for j in jobs] == ["Data Analyst", "Junior Analyst"]
    assert jobs[0].responsibilities == [
        "Built SQL dashboards used by 40 regional managers across the sales organization"
    ]


def test_two_line_resumes_keep_their_title_lines():
    text = (
        "Riley Park\n\nEXPERIENCE\nData Analyst\nAcme Corp | Denver, CO | 2021 - Present\n"
        "• Built dashboards\nJunior Analyst\nGlobex | Boston, MA | 2019 - 2020\n• Cleaned data\n"
        "\nEDUCATION\nBS Economics, 2018\n"
    )
    jobs = ResumeParser()._parse_text(text).work_experience
    assert [(j.title, j.employer) for j in jobs] == [("Data Analyst", "Acme Corp"), ("Junior Analyst", "Globex")]
    assert jobs[0].responsibilities == ["Built dashboards"]
