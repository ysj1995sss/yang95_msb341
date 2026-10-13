"""Spec 022: levels 2-3 reach a Word resume's skills line as their own reviewable change.
Synthetic documents and reviews; no model, no Word."""

import io

import pytest
from docx import Document

from resume_tailorer.analyzers.gap_analyzer import GapCategory, GapItem, GapReport
from resume_tailorer.analyzers.job_analyzer import JobAnalyzer
from resume_tailorer.analyzers.requirement_review import DIRECT, NONE, EvidenceRef, RequirementReview, RequirementRow
from resume_tailorer.artifacts.models import ChangeCategory
from resume_tailorer.docx_export import pipeline as pipeline_module
from resume_tailorer.docx_export.converter import DocxConversionUnavailable
from resume_tailorer.docx_export.pipeline import run_docx_tailoring_pipeline
from resume_tailorer.models import CareerTruthProfile, WorkExperience
from resume_tailorer.tailorer.docx_bullet_tailorer import BulletEdit, BulletTailoringResult

_BULLET = "Built Tableau dashboards used by 40 managers"


def _docx(skills_line="SQL, Excel"):
    doc = Document()
    for line in ("Sam Example", "sam@example.com | (555) 010-0100", "EXPERIENCE", "Data Analyst", "Acme | 2021 - 2024"):
        doc.add_paragraph(line)
    doc.add_paragraph(_BULLET, style="List Bullet")
    if skills_line is not None:
        doc.add_paragraph("SKILLS")
        doc.add_paragraph(skills_line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _review():
    evidence = EvidenceRef(_BULLET, "Data Analyst at Acme, bullet 1", "work_experience[0]", "experience", True)
    rows = (
        RequirementRow("r1", "Tableau dashboards", "required", False, DIRECT, ("Tableau",), (evidence,), "",
                       term_status=(("Tableau", DIRECT),)),
        RequirementRow("r2", "PostgreSQL", "preferred", False, NONE, ("PostgreSQL",), (), "",
                       term_status=(("PostgreSQL", NONE),)),
    )
    return RequirementReview(rows=rows, listed=("SQL", "Excel"))


class _KeepAll:
    seen: list = []

    def tailor_bullets(self, bullets, profile, job_analysis, gap_report, **kwargs):
        _KeepAll.seen = [b.section for b in bullets]
        return BulletTailoringResult(edits=[BulletEdit(b.paragraph_index, b.text, b.text, changed=False) for b in bullets],
                                     warnings=[])


@pytest.fixture
def no_word(monkeypatch):
    monkeypatch.setattr(pipeline_module, "convert_docx_to_pdf",
                        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("no Word in tests")))


def _profile():
    return CareerTruthProfile(contact_info={"name": "Sam Example", "email": "sam@example.com"}, education=[],
                              work_experience=[WorkExperience(employer="Acme", title="Data Analyst", dates="2021 - 2024",
                                                              responsibilities=[_BULLET], accomplishments=[])],
                              skills=["SQL", "Excel"], tools=[], certifications=[], accomplishments=[])


def _run(docx_bytes, review, gaps=None):
    return run_docx_tailoring_pipeline(docx_bytes, _profile(), JobAnalyzer().analyze("Tableau and PostgreSQL."),
                                       gaps or GapReport(items=[], summary=""), bullet_tailorer=_KeepAll(),
                                       convert_to_pdf=True, review=review)


def _gaps():
    """What a live gap report says about an unevidenced job term: truly missing, never add."""
    return GapReport(items=[GapItem(requirement="PostgreSQL", category=GapCategory.E,
                                    reason="Skill not found; do not add", candidate_evidence="None")], summary="")


def test_levels_2_and_3_are_one_reviewable_skills_change(no_word):
    result = _run(_docx(), _review())
    [change] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    assert change.section == "skills" and change.original_text == "SQL, Excel"
    assert change.proposed_text == "SQL, Excel, Tableau, PostgreSQL"
    assert [(term, level) for term, level, _reason in change.added_skills] == [("Tableau", 2), ("PostgreSQL", 3)]
    assert "skills" not in _KeepAll.seen  # the model never edits the skills line
    texts = [p.text for p in Document(io.BytesIO(result.docx_bytes)).paragraphs]
    assert "SQL, Excel, Tableau, PostgreSQL" in texts


def test_no_skills_line_or_no_review_adds_nothing(no_word):
    assert not [c for c in _run(_docx(skills_line=None), _review()).changes if c.category is ChangeCategory.SKILLS_ADDED]
    assert not [c for c in _run(_docx(), None).changes if c.category is ChangeCategory.SKILLS_ADDED]


def test_additions_that_push_the_resume_onto_another_page_are_dropped(no_word, monkeypatch):
    real = pipeline_module.finalize_docx_edits

    def finalize(**kwargs):
        result = real(**kwargs)
        grew = any(e.new_text.endswith("PostgreSQL") for e in kwargs["edits"])
        result.page_count_preserved = not grew
        return result

    monkeypatch.setattr(pipeline_module, "finalize_docx_edits", finalize)
    result = _run(_docx(), _review())
    # Spec 022 step 6: only what doesn't fit goes, lowest priority first, and the review says which.
    [change] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    assert [term for term, _level, _reason in change.added_skills] == ["Tableau"]
    assert change.proposed_text == "SQL, Excel, Tableau"
    assert result.skills_note == ("Left out of your skills line: PostgreSQL, because adding them pushed your "
                                  "resume onto another page.")


def test_additions_are_not_blamed_when_the_resume_was_already_too_long(no_word, monkeypatch):
    real = pipeline_module.finalize_docx_edits

    def finalize(**kwargs):
        result = real(**kwargs)
        result.page_count_preserved = False  # the rewrites alone already add a page
        return result

    monkeypatch.setattr(pipeline_module, "finalize_docx_edits", finalize)
    result = _run(_docx(), _review())
    assert not [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    assert result.skills_note == ("Left out of your skills line: Tableau, PostgreSQL, because your resume was "
                                  "already longer than the original.")


def _review_state(result, dispositions, manual_texts=None, gaps=None):
    from resume_tailorer.artifacts.models import FidelityMode

    return {"source_kind": "DOCX", "original_bytes": _docx(), "changes": list(result.changes),
            "dispositions": dispositions, "manual_texts": manual_texts or {}, "profile": _profile(),
            "requirement_review": _review(), "gap_report": gaps or GapReport(items=[], summary=""),
            "job_analysis": JobAnalyzer().analyze("Tableau and PostgreSQL."), "original_alignment": 0.5,
            "candidate_fit": None, "company": "Acme", "role": "Data Analyst",
            "fidelity_mode": FidelityMode.PRESERVED, "version": 1}


def test_an_accepted_addition_rebuilds_and_a_persons_own_edit_is_still_checked(no_word):
    from resume_tailorer.tailoring_service import TailoringError, regenerate

    result = _run(_docx(), _review())
    [change] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    state = _review_state(result, {change.change_id: "ACCEPTED"})
    regenerate({}, state)
    texts = [p.text for p in Document(io.BytesIO(state["docx_bytes"])).paragraphs]
    assert "SQL, Excel, Tableau, PostgreSQL" in texts

    edited = _review_state(result, {change.change_id: "MANUALLY_EDITED"},
                           {change.change_id: "SQL, Excel, Kubernetes"})
    with pytest.raises(TailoringError):
        regenerate({}, edited)


def test_a_computed_adjacent_skill_reaches_review_though_the_gap_report_calls_it_missing(no_word):
    """Spec 022 step 4: only the computed additions may pass the missing-requirement gate. Live gap
    reports list an unevidenced job term as truly missing, which used to reject the whole skills
    change (its level-2 terms too) before anyone saw it (code review, 2026-10-13)."""
    from resume_tailorer.ui.tailoring_view import group_changes, safe_default_dispositions

    result = _run(_docx(), _review(), gaps=_gaps())
    [change] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    assert change in group_changes(result.changes, ["PostgreSQL"]).reviewable
    assert change.change_id not in safe_default_dispositions(result.changes, ["PostgreSQL"])
    assert "PostgreSQL" in result.tailored_scoring_text  # scored as the file reads...
    assert "PostgreSQL" not in result_unsupported(result)  # ...and still not an unsupported claim


def result_unsupported(result):
    from resume_tailorer.analyzers.skill_additions import without_computed_skills
    from resume_tailorer.analyzers.gap_analyzer import find_unsupported_claims

    return find_unsupported_claims(without_computed_skills(_gaps(), result.changes, in_document=True),
                                   result.tailored_scoring_text)


def test_the_same_term_written_anywhere_else_is_still_blocked(no_word):
    """The exemption is the computed skills line only: a bullet, or a person's own skills wording,
    that adds the missing term is still kept out."""
    from dataclasses import replace

    from resume_tailorer.artifacts.models import ChangeDisposition
    from resume_tailorer.analyzers.skill_additions import without_computed_skills
    from resume_tailorer.ui.tailoring_view import group_changes

    result = _run(_docx(), _review(), gaps=_gaps())
    [skills] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    bullet = replace(skills, change_id="paragraph:5", section="work_experience", category=ChangeCategory.REPHRASED,
                     original_text=_BULLET, proposed_text=_BULLET + " on PostgreSQL", added_skills=())
    assert bullet in group_changes([skills, bullet], ["PostgreSQL"]).blocked
    edited = replace(skills, disposition=ChangeDisposition.MANUALLY_EDITED, manual_text="SQL, Excel, PostgreSQL")
    rejected = replace(skills, disposition=ChangeDisposition.REJECTED)
    for change in (edited, rejected):
        # A proposal is still judged as computed, but the finished document doesn't carry it.
        assert change in group_changes([change], ["PostgreSQL"]).reviewable
        assert [i.requirement for i in without_computed_skills(_gaps(), [change], in_document=True).items] == ["PostgreSQL"]


def test_an_accepted_adjacent_skill_is_not_reported_as_an_unsupported_claim(no_word):
    from resume_tailorer.tailoring_service import regenerate

    result = _run(_docx(), _review(), gaps=_gaps())
    [change] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    state = _review_state(result, {change.change_id: "ACCEPTED"}, gaps=_gaps())
    regenerate({}, state)
    assert "SQL, Excel, Tableau, PostgreSQL" in [p.text for p in Document(io.BytesIO(state["docx_bytes"])).paragraphs]
    assert "PostgreSQL" in state["tailored_text"] and "PostgreSQL" not in state["report"].unsupported_claims


def test_the_requirement_review_shows_an_added_adjacent_skill_as_added_not_missing(no_word):
    """Spec 022 step 5: the requirement review "shows the requirement as 'added to skills (adjacent)',
    not as evidence", and "Missing, never added" stops listing a term that was added."""
    from dataclasses import replace

    from resume_tailorer.artifacts.models import ChangeDisposition
    from resume_tailorer.ui.requirement_review_view import adjacent_additions, missing_list, review_view

    result = _run(_docx(), _review(), gaps=_gaps())
    adjacent = adjacent_additions(result.changes)
    assert adjacent == {"postgresql"}  # Tableau is level 2: the review already shows its evidence
    rows = {row["text"]: row for group in review_view(_review(), adjacent)["groups"] for row in group["rows"]}
    assert rows["PostgreSQL"]["label"] == "Added to skills (adjacent)"
    assert rows["PostgreSQL"]["terms"][0]["label"] == "Added to skills (adjacent)"
    assert rows["Tableau dashboards"]["status"] == "direct"
    assert missing_list(_review(), adjacent) == []
    assert [m["label"] for m in missing_list(_review())] == ["PostgreSQL"]
    # Kept original: the file no longer names it, so it's missing again.
    assert adjacent_additions([replace(c, disposition=ChangeDisposition.REJECTED) for c in result.changes]) == set()


def test_a_length_correction_keeps_the_changes_the_file_still_contains(no_word, monkeypatch):
    """The correction rebuild paired bare edits again, so reviewed changes lost their reasons and a
    skills change vanished from the review while the file kept its added skills (2026-10-13). The
    skills line is never reworded to fit: additions are dropped, not condensed (spec 022 step 6)."""
    from resume_tailorer.artifacts import length_control

    result = _run(_docx(), _review())
    monkeypatch.setattr(length_control, "needs_length_correction", lambda validation: True)
    corrected = length_control.correct_docx_length_once(original_docx_bytes=_docx(), docx_result=result,
                                                         profile=_profile(), gap_report=GapReport(items=[], summary=""))
    [change] = [c for c in corrected.changes if c.category is ChangeCategory.SKILLS_ADDED]
    assert change.proposed_text == "SQL, Excel, Tableau, PostgreSQL" and change.reason == _skills(result).reason
    assert "SQL, Excel, Tableau, PostgreSQL" in [p.text for p in Document(io.BytesIO(corrected.docx_bytes)).paragraphs]


def _skills(result):
    return next(c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED)


def test_a_refused_edit_says_once_in_plain_words_what_it_adds(no_word):
    """Dogfood, 2026-10-13: adding Terraform showed "⚠️ FABRICATION RISK: 'terraform' mentioned in …",
    then the same thing twice more in other words, calling a skills line "this project"."""
    from resume_tailorer.tailoring_service import TailoringError, regenerate

    result = _run(_docx(), _review())
    [change] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    edited = _review_state(result, {change.change_id: "MANUALLY_EDITED"}, {change.change_id: "SQL, Excel, Terraform"})
    with pytest.raises(TailoringError) as refused:
        regenerate({}, edited)
    message = str(refused.value)
    assert "Terraform" in message and message.lower().count("terraform") == 1
    for jargon in ("FABRICATION", "⚠", "project", "mentioned in"):
        assert jargon not in message
    assert "Career Profile" in message


def test_a_longer_but_truthful_edit_is_not_refused_for_its_length(no_word):
    """Length is a soft warning for model rewrites; it used to refuse a person's own edit outright."""
    from resume_tailorer.tailoring_service import regenerate

    result = _run(_docx(), _review())
    [change] = [c for c in result.changes if c.category is ChangeCategory.SKILLS_ADDED]
    state = _review_state(result, {change.change_id: "MANUALLY_EDITED"}, {change.change_id: "SQL, Excel, Tableau"})
    regenerate({}, state)
    assert "SQL, Excel, Tableau" in [p.text for p in Document(io.BytesIO(state["docx_bytes"])).paragraphs]


def _competencies_docx(model_line=None):
    doc = Document()
    for line in ("Sam Example", "sam@example.com | (555) 010-0100", "EXPERIENCE", "Data Analyst", "Acme | 2021 - 2024"):
        doc.add_paragraph(line)
    doc.add_paragraph(_BULLET, style="List Bullet")
    doc.add_paragraph("ADDITIONAL")
    doc.add_paragraph("Core Competencies: SQL | Excel", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class _ReordersCompetencies(_KeepAll):
    def tailor_bullets(self, bullets, profile, job_analysis, gap_report, **kwargs):
        result = super().tailor_bullets(bullets, profile, job_analysis, gap_report, **kwargs)
        for edit in result.edits:
            if edit.original_text.startswith("Core Competencies"):
                edit.new_text, edit.changed = "Core Competencies: Excel | SQL", True
        return result


@pytest.mark.parametrize("tailorer, line", [(_KeepAll(), "Core Competencies: SQL | Excel | Tableau | PostgreSQL"),
                                            (_ReordersCompetencies(), "Core Competencies: Excel | SQL | Tableau | PostgreSQL")])
def test_a_core_competencies_line_takes_the_additions_when_there_is_no_skills_line(no_word, tailorer, line):
    """Spec 022 step 3 says "the existing skills/competencies paragraph", but a "Core Competencies:"
    line was skipped, so resumes whose only list is that line never got additions (code review,
    2026-10-13). The model's reorder of the line is kept, with the additions after it, as one change."""
    result = run_docx_tailoring_pipeline(_competencies_docx(), _profile(), JobAnalyzer().analyze("Tableau and PostgreSQL."),
                                         GapReport(items=[], summary=""), bullet_tailorer=tailorer,
                                         convert_to_pdf=True, review=_review())
    lines = [c for c in result.changes if c.original_text.startswith("Core Competencies")]
    assert len(lines) == 1 and lines[0].category is ChangeCategory.SKILLS_ADDED and lines[0].proposed_text == line
    assert line in [p.text for p in Document(io.BytesIO(result.docx_bytes)).paragraphs]
    assert "PostgreSQL" in result.tailored_scoring_text  # the review reads the line as the file has it


def test_a_reordered_competencies_line_does_not_replace_the_profile_skills_in_scoring(no_word):
    """Code review, 2026-10-13: with no skills line, a model reorder of "Core Competencies:" became the
    scoring text's whole skills list, so skills only in the profile dropped out of the keyword report."""
    from dataclasses import replace as replace_profile

    person = replace_profile(_profile(), skills=["SQL", "Excel", "Python"])
    result = run_docx_tailoring_pipeline(_competencies_docx(), person, JobAnalyzer().analyze("Python and SQL."),
                                         GapReport(items=[], summary=""), bullet_tailorer=_ReordersCompetencies(),
                                         convert_to_pdf=True, review=None)
    assert "Python" in result.tailored_scoring_text
    assert "Core Competencies: Excel | SQL" in result.tailored_scoring_text
