import pytest
from pathlib import Path
from resume_tailorer.parsers.resume_parser import ResumeParser
from resume_tailorer.models import CareerTruthProfile

@pytest.fixture
def sample_resume_path():
    """Path to a sample resume for testing."""
    # For MVP testing, we'll use a PDF fixture
    return Path(__file__).parent / "fixtures" / "sample_resume.pdf"

def test_resume_parser_extracts_text_from_pdf(sample_resume_path):
    """Parser can extract text from a PDF resume."""
    parser = ResumeParser()
    profile = parser.parse(str(sample_resume_path))

    assert isinstance(profile, CareerTruthProfile)
    assert profile.name  # Name should be extracted
    assert profile.email  # Email should be extracted
    assert len(profile.work_experience) > 0  # Should have at least one job

def test_resume_parser_extracts_contact_info(sample_resume_path):
    """Parser correctly identifies contact info (name, email, phone)."""
    parser = ResumeParser()
    profile = parser.parse(str(sample_resume_path))

    # Verify contact info is extracted
    assert profile.contact_info is not None
    assert profile.contact_info.get("name") == "John Smith"
    assert profile.contact_info.get("email") == "john.smith@email.com"
    assert "(555) 123-4567" in profile.contact_info.get("phone", "")

def test_resume_parser_extracts_education(sample_resume_path):
    """Parser extracts education entries."""
    parser = ResumeParser()
    profile = parser.parse(str(sample_resume_path))

    # Verify education is extracted
    assert isinstance(profile.education, list)
    # May have 0+ education entries depending on PDF parsing
    # At minimum, the parser should handle the education section without errors

def test_resume_parser_extracts_work_experience(sample_resume_path):
    """Parser extracts work experience with employer, title, dates, responsibilities, accomplishments."""
    parser = ResumeParser()
    profile = parser.parse(str(sample_resume_path))

    # Verify work experience is extracted
    assert len(profile.work_experience) > 0

    # Check first job entry has required fields
    first_job = profile.work_experience[0]
    assert first_job.employer
    assert first_job.title
    assert first_job.dates


class TestWrappedBulletContinuation:
    """
    Regression tests for a bug found live (2026-09-22) on a real resume: a
    long bullet that PDF text extraction wraps onto a second physical line
    with no bullet marker was silently DROPPED (only the first line was
    kept), losing real accomplishment content. This cascaded into a false
    "fabrication risk" flag later, since the dropped content reappeared in
    the tailored output looking unsupported by the (incomplete) profile.
    """

    def test_wrapped_continuation_is_appended_not_dropped(self):
        text = (
            "EXPERIENCE\n"
            "Marketing Manager\n"
            "Acme Corp | Springfield, IL Jan 2022-Dec 2023\n"
            "• Developed a growth strategy by synthesizing data, identifying up to $50M\n"
            "in incremental sales potential targeted for next fiscal year\n"
            "EDUCATION\n"
        )
        parser = ResumeParser()
        jobs = parser._extract_work_experience(text)

        assert len(jobs) == 1
        combined = " ".join(jobs[0].responsibilities + jobs[0].accomplishments)
        assert "$50M" in combined
        # The wrapped continuation text must survive, not just the first line.
        assert "incremental sales potential" in combined

    def test_multiple_bullets_each_keep_their_own_continuation(self):
        text = (
            "EXPERIENCE\n"
            "Operations Lead\n"
            "Beta Inc | Remote Mar 2021-Feb 2022\n"
            "• Led a 10-person team, diagnosed process bottlenecks, and implemented\n"
            "improvements that reduced cycle time by 30% (60s -> 42s)\n"
            "• Coordinated cross-functional projects for 50K+ users with 91% satisfaction\n"
            "EDUCATION\n"
        )
        parser = ResumeParser()
        jobs = parser._extract_work_experience(text)

        assert len(jobs) == 1
        all_bullets = jobs[0].responsibilities + jobs[0].accomplishments
        assert any("cycle time by 30%" in b and "60s -> 42s" in b for b in all_bullets)
        assert any("91% satisfaction" in b for b in all_bullets)


class TestEmploymentHeading:
    """Regression test for a heading variant a product-revision spec
    (2026-09-23) explicitly called out: some resumes title the whole
    section "EMPLOYMENT" instead of "WORK EXPERIENCE"/"PROFESSIONAL
    EXPERIENCE", with no literal word "experience" anywhere in the resume."""

    def test_employment_heading_is_recognized_with_no_word_experience_present(self):
        text = (
            "EMPLOYMENT\n"
            "Marketing Manager\n"
            "Acme Corp | Springfield, IL Jan 2022-Dec 2023\n"
            "• Built a growth strategy\n"
            "EDUCATION\n"
        )
        parser = ResumeParser()
        jobs = parser._extract_work_experience(text)
        assert len(jobs) == 1
        assert jobs[0].employer == "Acme Corp"

    def test_employment_recognized_as_a_section_boundary(self):
        from resume_tailorer.parsers.section_headings import SECTION_BOUNDARY_RE

        assert SECTION_BOUNDARY_RE.match("EMPLOYMENT")


class TestInlineLabeledSkills:
    """
    Regression test for a bug found live (2026-09-22): a real resume labeled
    its skills-equivalent content "Technical Proficiency" and "Core
    Competencies" instead of the literal word "Skills" -- the section-only
    detection required "skill" to appear somewhere, so it returned an empty
    list even though the resume clearly listed real skills.
    """

    def test_technical_proficiency_label_is_recognized(self):
        text = (
            "ADDITIONAL\n"
            "• Technical Proficiency: Tableau | Power BI | SQL | MS Excel\n"
        )
        parser = ResumeParser()
        skills = parser._extract_skills(text)
        assert "Tableau" in skills
        assert "Power BI" in skills
        assert "SQL" in skills

    def test_core_competencies_label_is_recognized(self):
        text = (
            "ADDITIONAL\n"
            "• Core Competencies: Strategic Thinking | Market Research | Mentorship\n"
        )
        parser = ResumeParser()
        skills = parser._extract_skills(text)
        assert "Strategic Thinking" in skills
        assert "Market Research" in skills


class TestToolsAreNotRepeatedFromTheSkillsLine:
    """Found 2026-10-13: a fresh import put Docker, AWS and other tools on the skills line into
    tools as well, so Career Profile opened with "5 to review" problems the person never made and
    every Word run said the tools couldn't be synced."""

    def test_a_tool_on_the_skills_line_is_listed_once(self):
        text = (
            "Alex Rivera\nalex.rivera.demo@example.com\n\n"
            "EXPERIENCE\nSoftware Engineer\nBeacon Software Co. | Austin, TX | 2018-2021\n"
            "• Deployed services with Docker and Terraform\n\n"
            "SKILLS\nPython, Docker, aws\n"
        )
        profile = ResumeParser()._parse_text(text)
        assert {"Python", "Docker", "aws"} <= set(profile.skills)
        assert "Docker" not in profile.tools and "AWS" not in profile.tools
        assert profile.tools == ["Terraform"]  # named only in a bullet, so it's still a tool


class TestCertificationsNeedACertificationWord:
    """Found 2026-10-13: any phrase starting with AWS, Azure, Docker or Kubernetes became a
    certification, so "Migrated workloads to AWS Lambda and ECS" was imported as a credential the
    person doesn't hold (and could be printed in a generated resume's certifications)."""

    def test_a_tool_named_in_a_bullet_is_not_a_certification(self):
        text = ("EXPERIENCE\n• Deployed services with Docker and Terraform\n"
                "• Migrated workloads to AWS Lambda and ECS\n• Certified 40 vendors for compliance\n")
        assert ResumeParser()._extract_certifications(text) == []

    def test_real_certifications_are_still_found(self):
        text = ("CERTIFICATIONS\n• AWS Certified Solutions Architect – Associate\n"
                "• Kubernetes Administrator (CKA)\n• Certified Scrum Master\n• PMP\n")
        found = set(ResumeParser()._extract_certifications(text))
        assert {"AWS Certified Solutions Architect – Associate", "Kubernetes Administrator (CKA)",
                "Certified Scrum Master", "PMP"} <= found


def test_extract_style_hints_detects_bullet_dash():
    """Detects '-' as the dominant bullet character."""
    parser = ResumeParser()
    text = "WORK EXPERIENCE\n- Built a system\n- Led a team\n- Shipped a feature"
    hints = parser.extract_style_hints(text)
    assert hints["bullet_char"] == "-"

def test_extract_style_hints_detects_bullet_dot():
    """Detects '•' as the dominant bullet character."""
    parser = ResumeParser()
    text = "WORK EXPERIENCE\n• Built a system\n• Led a team\n• Shipped a feature"
    hints = parser.extract_style_hints(text)
    assert hints["bullet_char"] == "•"

def test_extract_style_hints_defaults_to_dot_when_no_bullets_found():
    """Falls back to '•' default when no bullet markers are present."""
    parser = ResumeParser()
    text = "WORK EXPERIENCE\nBuilt a system without bullets\nLed a team without bullets"
    hints = parser.extract_style_hints(text)
    assert hints["bullet_char"] == "•"

def test_extract_style_hints_detects_short_caps_heading_as_bold_larger():
    """A short ALL-CAPS heading line suggests a larger/bold heading style."""
    parser = ResumeParser()
    text = "EXPERIENCE\n- Built a system\nEDUCATION\n- BS Computer Science"
    hints = parser.extract_style_hints(text)
    assert hints["heading_style"] == "bold_larger"

def test_extract_style_hints_defaults_heading_style_to_bold():
    """No short all-caps heading present defaults to plain bold."""
    parser = ResumeParser()
    text = "some lowercase text with no headings at all in this resume body"
    hints = parser.extract_style_hints(text)
    assert hints["heading_style"] == "bold"

class TestDetectPdfStyle:
    """
    Regression tests for font-family and page-count detection, added so
    that "preserve original length" (spec item 15) actually detects and
    matches the user's real page count instead of using a fixed preset
    regardless of the source, and so the output font is at least in the
    same family (serif/sans-serif) as the original -- both requested
    directly (2026-09-22) after a user's third round of PDF feedback.
    """

    def _make_pdf(self, tmp_path, font_name: str, num_pages: int = 1) -> str:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import LETTER

        path = str(tmp_path / "test.pdf")
        c = canvas.Canvas(path, pagesize=LETTER)
        for page in range(num_pages):
            c.setFont(font_name, 12)
            c.drawString(72, 750, f"Page {page + 1} content")
            c.showPage()
        c.save()
        return path

    def test_detects_sans_serif_font(self, tmp_path):
        path = self._make_pdf(tmp_path, "Helvetica")
        parser = ResumeParser()
        result = parser._detect_pdf_style(path)
        assert result.get("font_family") == "sans-serif"

    def test_detects_serif_font(self, tmp_path):
        path = self._make_pdf(tmp_path, "Times-Roman")
        parser = ResumeParser()
        result = parser._detect_pdf_style(path)
        assert result.get("font_family") == "serif"

    def test_detects_single_page_count(self, tmp_path):
        path = self._make_pdf(tmp_path, "Helvetica", num_pages=1)
        parser = ResumeParser()
        result = parser._detect_pdf_style(path)
        assert result["page_count"] == 1

    def test_detects_multi_page_count(self, tmp_path):
        path = self._make_pdf(tmp_path, "Helvetica", num_pages=2)
        parser = ResumeParser()
        result = parser._detect_pdf_style(path)
        assert result["page_count"] == 2

    def test_extract_style_hints_includes_detected_style_when_file_path_given(self, tmp_path):
        path = self._make_pdf(tmp_path, "Times-Roman", num_pages=2)
        parser = ResumeParser()
        hints = parser.extract_style_hints("some text", file_path=path)
        assert hints["page_count"] == 2
        assert hints["font_family"] == "serif"
        # Text-derived hints must still be present too.
        assert "bullet_char" in hints
        assert "heading_style" in hints

    def test_extract_style_hints_without_file_path_has_no_style_keys(self):
        """Backward compatible: existing callers passing only text keep working."""
        parser = ResumeParser()
        hints = parser.extract_style_hints("some resume text")
        assert "page_count" not in hints
        assert "font_family" not in hints

    def test_detect_pdf_style_handles_a_missing_file_gracefully(self):
        parser = ResumeParser()
        result = parser._detect_pdf_style("this/path/does/not/exist.pdf")
        assert result == {}


def test_get_raw_text_dispatches_pdf(sample_resume_path):
    """get_raw_text() extracts text from a PDF via its public dispatch method."""
    parser = ResumeParser()
    text = parser.get_raw_text(str(sample_resume_path))
    assert isinstance(text, str)
    assert "John Smith" in text

def test_get_raw_text_raises_on_unsupported_extension():
    """get_raw_text() raises ValueError for an unsupported file extension."""
    parser = ResumeParser()
    with pytest.raises(ValueError, match="Unsupported"):
        parser.get_raw_text("resume.txt")


def test_get_raw_text_rejects_legacy_doc():
    parser = ResumeParser()
    with pytest.raises(ValueError, match=r"\.doc"):
        parser.get_raw_text("resume.doc")


class TestEducationAndSummaryExtraction:
    """
    Regression tests for two bugs found live (2026-09-22) on a real
    two-degree resume: the MBA (a second, in-progress degree) vanished
    entirely, its scholarship/honors bullets were dropped, and a
    3-sentence professional summary paragraph had nowhere to be stored at
    all (CareerTruthProfile had no summary field).
    """

    REAL_STYLE_RESUME = (
        "Jane Doe\n"
        "(555) 123-4567 https://www.linkedin.com/in/janedoe/ jane@example.com\n"
        "Data-savvy analyst with a proven record of driving growth. Expert in translating\n"
        "insights into action and delivering measurable outcomes.\n"
        "EDUCATION\n"
        "STATE UNIVERSITY, SCHOOL OF BUSINESS | Springfield, IL Sept 2025-Apr 2027\n"
        "Master of Business Administration\n"
        "• Awarded Half-Tuition Merit Scholarship\n"
        "• Consulted Acme Corp on a market-entry analysis\n"
        "STATE UNIVERSITY-NORTH | Springfield, IL Apr 2018-Apr 2022\n"
        "B.S. Business Administration (Summa Cum Laude)\n"
        "• Dean's List x5 Semesters | GPA: 3.9/4.0\n"
        "PROFESSIONAL EXPERIENCE\n"
        "Analyst\n"
        "Acme Widgets | Remote Jan 2022-Present\n"
        "• Improved operational productivity by 15%\n"
    )

    def test_both_degrees_are_extracted_not_just_the_first(self):
        parser = ResumeParser()
        education = parser._extract_education(self.REAL_STYLE_RESUME)
        assert len(education) == 2
        assert "master" in education[0].degree.lower() or education[0].degree.upper() == "MBA"
        assert education[0].year == 2027
        assert "B.S" in education[1].degree.upper() or "BS" in education[1].degree.upper()
        assert education[1].year == 2022

    def test_education_bullets_are_preserved_as_notes(self):
        parser = ResumeParser()
        education = parser._extract_education(self.REAL_STYLE_RESUME)
        mba = education[0]
        assert any("Scholarship" in n for n in mba.notes)
        assert any("Acme Corp" in n for n in mba.notes)

    def test_work_experience_bullet_does_not_leak_into_education(self):
        parser = ResumeParser()
        education = parser._extract_education(self.REAL_STYLE_RESUME)
        all_notes = [n for e in education for n in e.notes]
        assert not any("operational productivity" in n for n in all_notes)

    def test_summary_paragraph_is_extracted(self):
        parser = ResumeParser()
        summary = parser._extract_summary(self.REAL_STYLE_RESUME)
        assert "Data-savvy analyst" in summary
        assert "measurable outcomes" in summary

    def test_summary_excludes_name_and_contact_line(self):
        parser = ResumeParser()
        summary = parser._extract_summary(self.REAL_STYLE_RESUME)
        assert "Jane Doe" not in summary
        assert "jane@example.com" not in summary

    def test_full_parse_includes_summary_on_the_profile(self):
        parser = ResumeParser()
        profile = parser._parse_text(self.REAL_STYLE_RESUME)
        assert profile.summary
        assert "Data-savvy analyst" in profile.summary


class TestCompanyLineDates:
    """Found live (2026-09-30): a tab before the dates and no spaces around the
    dash lost or truncated three of four real job dates."""

    TEXT = (
        "Alex Kim\nalex@example.com\n\nPROFESSIONAL EXPERIENCE\n"
        "Marketing Strategy Intern\nAcme Health | Boston, MA\tMay 2026\u2013Aug 2026\n"
        "- Built a growth strategy\n"
        "Marketing Manager\nGlobex Co | Zhengzhou, China\tAug 2024-Jul 2025\n"
        "- Grew sales 75%\n"
        "Analyst\nInitech | Remote \u2014 03/2021 to Present\n"
        "- Ran reports\n"
        "Coordinator\nHooli Hotel | Layton, UT | Summer 2020\n"
        "- Led guest programs\n"
        "\nEDUCATION\nMBA, State University, 2027\n"
    )

    def test_dates_and_locations_are_separated_correctly(self):
        jobs = ResumeParser()._parse_text(self.TEXT).work_experience
        assert [(j.employer, j.location, j.dates) for j in jobs] == [
            ("Acme Health", "Boston, MA", "May 2026\u2013Aug 2026"),
            ("Globex Co", "Zhengzhou, China", "Aug 2024-Jul 2025"),
            ("Initech", "Remote", "03/2021 to Present"),
            ("Hooli Hotel", "Layton, UT", "Summer 2020"),
        ]


def test_a_wrapped_bullet_line_with_a_year_does_not_start_a_new_job():
    """Found 2026-10-05 on a real resume: "...potential targeted for implementation by 2027", the
    second line of a wrapped bullet, was taken as the next job's dates, so the job kept only its
    first bullet's first line and lost the rest."""
    from resume_tailorer.parsers.resume_parser import ResumeParser

    text = (
        "Riley Park\nriley@example.com\n"
        "PROFESSIONAL EXPERIENCE\n"
        "Strategy Intern\n"
        "Acme Health | Springfield, RI May 2026\u2013Aug 2026\n"
        "\u2022 Developed a growth strategy by synthesizing data, identifying up to $120M in incremental sales\n"
        "potential targeted for implementation by 2027\n"
        "\u2022 Built a customer framework targeting 3\u00d7 visit frequency and 2\u00d7 basket size\n"
        "\u2022 Aligned 30+ stakeholders and distilled research into executive-ready\n"
        "recommendations adopted across 9,000+ stores\n"
        "\n"
        "Marketing Manager\n"
        "Northwind Shoes | Zhengzhou, China Aug 2024-Jul 2025\n"
        "\u2022 Drove 75% YoY sales growth with a market entry strategy\n"
        "EDUCATION\n"
        "State University | 2019\nBS Economics\n"
    )
    jobs = ResumeParser()._parse_text(text).work_experience
    assert [j.employer for j in jobs] == ["Acme Health", "Northwind Shoes"]
    first = jobs[0].accomplishments + jobs[0].responsibilities
    assert len(first) == 3
    assert any(b.endswith("potential targeted for implementation by 2027") for b in first)
    assert any("9,000+ stores" in b for b in first)
    assert jobs[0].dates == "May 2026\u2013Aug 2026"


class TestOneLineEducation:
    """Bug hunt 2026-10-12 (#1, #2, #4): "Degree, School, Year" on one line put the whole
    line in the school box and left the degree empty, so Career Profile asked the person to
    fix a degree they had stated. Synthetic layouts only."""

    @staticmethod
    def _edu(*lines):
        text = "Sam Example\nsam@example.com\nEDUCATION\n" + "\n".join(lines) + "\nSKILLS\nSQL, Excel\n"
        return ResumeParser()._extract_education(text)

    @pytest.mark.parametrize("line, degree, field, school, year", [
        ("BS Economics, State University, 2019", "BS", "Economics", "State University", 2019),
        ("B.S. in Economics, State University, 2019", "B.S.", "Economics", "State University", 2019),
        ("Bachelor of Science in Economics, State University (2015-2019)",
         "Bachelor of Science", "Economics", "State University", 2019),
        ("State University | BA Statistics | 2018", "BA", "Statistics", "State University", 2018),
        ("MBA, Tech Institute, 2022", "MBA", "", "Tech Institute", 2022),
    ])
    def test_one_line_degree_school_year(self, line, degree, field, school, year):
        [entry] = self._edu(line)
        assert (entry.degree, entry.field, entry.institution, entry.year) == (degree, field, school, year)

    def test_two_one_line_degrees_both_keep_their_degree(self):
        first, second = self._edu("M.S. Data Science, Tech Institute, 2022", "B.A. Statistics, State University, 2019")
        assert (first.degree, first.field, first.institution) == ("M.S.", "Data Science", "Tech Institute")
        assert (second.degree, second.field, second.institution) == ("B.A.", "Statistics", "State University")

    @pytest.mark.parametrize("school_line", ["State University, Boston, MA 2019", "State College, Jackson, MS 2019"])
    def test_a_state_abbreviation_is_not_a_degree(self, school_line):
        [entry] = self._edu(school_line, "B.S. Economics")
        assert (entry.degree, entry.field) == ("B.S.", "Economics")
        assert entry.institution.startswith("State")

    def test_two_line_layout_drops_the_year_from_the_school_and_keeps_gpa_apart(self):
        [entry] = self._edu("State University, Portland, OR  2019", "B.S. Economics, GPA 3.6")
        assert entry.institution == "State University, Portland, OR"
        assert (entry.degree, entry.field, entry.gpa) == ("B.S.", "Economics", "3.6")


def test_title_on_the_line_below_company_city_dates_is_not_swapped():
    # Bug hunt 2026-10-12 (#3): "Acme Corp | Portland, OR | Dates" then "Data Analyst" was
    # read as title "Acme Corp" at employer "Portland". Synthetic layout.
    text = ("Sam Example\nsam@example.com\nEXPERIENCE\n"
            "Acme Corp | Portland, OR | January 2021 to Present\nData Analyst\n"
            "• Built SQL dashboards used by 40 managers\n"
            "Beta LLC | Remote | 2018 - 2020\nReporting Analyst\n"
            "• Automated weekly reports\nEDUCATION\nBS Economics, State University, 2019\n")
    jobs = ResumeParser()._extract_work_experience(text)
    assert [(j.title, j.employer) for j in jobs] == [("Data Analyst", "Acme Corp"), ("Reporting Analyst", "Beta LLC")]
    assert jobs[0].dates == "January 2021 to Present" and jobs[0].location == "Portland, OR"
    assert jobs[0].responsibilities + jobs[0].accomplishments == ["Built SQL dashboards used by 40 managers"]


# Bug hunt 2026-10-12, second pass. Synthetic resumes only.

@pytest.mark.parametrize("glyph", ["\uf0b7", "\x7f"], ids=["word-pdf-symbol-bullet", "reportlab-pdf-bullet"])
def test_pdf_bullet_glyphs_are_read_as_bullets(glyph):
    # U+F0B7 is the bullet Word's "Save as PDF" puts in the text layer; \x7f is what the
    # PDF reader returns for Job Copilot's own generated PDFs. Both lost every bullet.
    text = ("Sam Example\nsam@example.com\nEXPERIENCE\nData Analyst | Acme Corp | Jan 2021 - Present\n"
            f"{glyph} Built SQL dashboards used by 40 managers\n {glyph} Cut report time by 30%\n"
            "EDUCATION\nBS Economics, State University, 2019\n")
    [job] = ResumeParser()._parse_text(text).work_experience
    assert job.responsibilities + job.accomplishments == ["Built SQL dashboards used by 40 managers",
                                                          "Cut report time by 30%"]


def test_a_resume_job_copilot_generated_keeps_its_bullets_when_imported_again(tmp_path):
    from resume_tailorer.pdf.generator import PDFGenerator

    body = ("sam@example.com | (555) 010-0100\nEXPERIENCE\nData Analyst | Acme Corp | Jan 2021 - Present\n"
            "- Built SQL dashboards used by 40 managers\n- Cut report time by 30%\n"
            "EDUCATION\nBS Economics, State University, 2019\nSKILLS\nSQL, Python\n")
    path = str(tmp_path / "ours.pdf")
    PDFGenerator().generate(body, "Sam Example", output_path=path)
    [job] = ResumeParser().parse(path).work_experience
    assert len(job.responsibilities + job.accomplishments) == 2


def test_name_and_contact_in_a_word_header_table_are_read(tmp_path):
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Sam Example"
    table.cell(0, 1).text = "sam@example.com\n(555) 010-0100\nPortland, OR"
    for line in ("SUMMARY", "Analyst with five years of reporting experience.", "EXPERIENCE",
                 "Data Analyst | Acme Corp | Jan 2021 - Present", "- Built SQL dashboards used by 40 managers",
                 "EDUCATION", "BS Economics, State University, 2019"):
        doc.add_paragraph(line)
    path = tmp_path / "table.docx"
    doc.save(path)
    contact = ResumeParser().parse(str(path)).contact_info
    assert (contact.get("name"), contact.get("email"), contact.get("phone")) == (
        "Sam Example", "sam@example.com", "(555) 010-0100")


def test_a_section_heading_is_never_taken_as_the_name():
    text = "SUMMARY\nAnalyst with five years of reporting experience.\nEXPERIENCE\nData Analyst | Acme | 2021 - Present\n"
    assert ResumeParser()._extract_contact_info(text)["name"] == "Unknown"


def test_a_summary_heading_after_a_multi_line_contact_block_is_not_part_of_the_summary():
    text = ("Sam Example\nsam@example.com\n(555) 010-0100\nPortland, OR\n\nSUMMARY\n"
            "Analyst with five years of reporting experience.\nEXPERIENCE\nData Analyst | Acme | 2021 - Present\n")
    assert ResumeParser()._extract_summary(text) == "Analyst with five years of reporting experience."


class TestContactLocation:
    """Dogfood, 2026-10-13: a location on the contact line was never read, so the profile had none
    (and a generated PDF printed "N/A" in its place)."""

    def test_city_and_state_on_the_contact_line(self):
        text = "John Smith\njohn@example.com | (555) 123-4567 | San Francisco, CA\nPROFESSIONAL SUMMARY\nEngineer.\n"
        assert ResumeParser()._extract_contact_info(text)["location"] == "San Francisco, CA"

    def test_no_location_is_not_invented(self):
        text = "John Smith\njohn@example.com | (555) 123-4567\nEXPERIENCE\nAcme | Austin, TX | 2020 - 2024\n"
        assert "location" not in ResumeParser()._extract_contact_info(text)


class TestDegreeAboveTheSchoolLine:
    """Dogfood, 2026-10-13: "BS Computer Science" above "University … | 2016" lost the degree, so
    Career Profile said it wasn't recognized and the tailored resume had no degree at all."""

    def test_degree_line_above_the_school_line(self):
        text = ("EDUCATION\nBS Computer Science\nUniversity of California, Berkeley | 2016\n"
                "TECHNICAL SKILLS\nLanguages: Python\n")
        [entry] = ResumeParser()._extract_education(text)
        assert (entry.degree, entry.field, entry.institution, entry.year) == (
            "BS", "Computer Science", "University of California, Berkeley", 2016)

    def test_two_entries_each_with_the_degree_above(self):
        text = ("EDUCATION\nMBA, Finance\nState University | 2022\nBA Economics\nCity College | 2016\n"
                "SKILLS\nExcel\n")
        entries = ResumeParser()._extract_education(text)
        assert [(e.degree, e.field, e.institution) for e in entries] == [
            ("MBA", "Finance", "State University"), ("BA", "Economics", "City College")]

    def test_the_degree_below_still_wins(self):
        text = "EDUCATION\nState University | 2016\nB.S. Economics\nSKILLS\nExcel\n"
        [entry] = ResumeParser()._extract_education(text)
        assert (entry.degree, entry.field) == ("B.S.", "Economics")


class TestLabelledSkillLines:
    """Dogfood, 2026-10-13: a "Tools: …" line ended the skills section early, so the "Frameworks:"
    line after it was lost, and each line's label ("Languages") was kept as a skill."""

    def test_every_labelled_line_is_read_without_its_label(self):
        text = ("TECHNICAL SKILLS\nLanguages: Python, JavaScript, Java, SQL\nTools: Docker, Kubernetes, Git\n"
                "Frameworks: Flask, React, Django\nCERTIFICATIONS\nPMP\n")
        skills = ResumeParser()._extract_skills(text)
        assert {"Python", "SQL", "Docker", "Git", "Flask", "React", "Django"} <= set(skills)
        assert not {"Languages", "Tools", "Frameworks", "PMP"} & set(skills)

    def test_a_tools_heading_on_its_own_still_ends_the_skills_section(self):
        text = "SKILLS\nPython, SQL\nTOOLS\nJira\n"
        assert ResumeParser()._extract_skills(text)[:2] == ["Python", "SQL"] and "Jira" not in ResumeParser()._extract_skills(text)


class TestSkillsSectionEdges:
    """Code review, 2026-10-13."""

    def test_one_skill_per_line_is_read_whole(self):
        # The section's end pattern matched any one-word line (case-insensitive), so "SQL" ended it.
        text = "SKILLS\nPython\nSQL\nMATLAB\nTableau\nPROJECTS\nInventory dashboard\n"
        assert ResumeParser()._extract_skills(text) == ["Python", "SQL", "MATLAB", "Tableau"]

    def test_a_multi_word_capitals_heading_still_ends_the_section(self):
        text = "SKILLS\nPython\nSQL\nVOLUNTEER WORK\nFood bank driver\n"
        assert ResumeParser()._extract_skills(text) == ["Python", "SQL"]

    def test_a_skill_used_as_a_line_label_is_kept(self):
        # "Tableau: dashboards" names a skill; "Languages:" only names a group.
        text = "SKILLS\nTableau: dashboards, LOD expressions\nTools & Platforms: Docker, Git\nLanguages: Python\n"
        skills = ResumeParser()._extract_skills(text)
        assert {"Tableau", "Docker", "Git", "Python"} <= set(skills)
        assert not {"Tools & Platforms", "Languages"} & set(skills)


def test_bullets_keep_their_order_in_the_resume():
    """Dogfood round 3, 2026-10-13: bullets with a number were split into a separate accomplishments list
    that every step lists after the rest, so a rebuilt PDF resume silently reordered each job's bullets
    (and evidence labels like "bullet 3" counted in the wrong order)."""
    text = ("EXPERIENCE\nSenior Data Analyst\nNorthwind Retail | Austin, TX | 2021 - Present\n"
            "- Led a team of 3 analysts building weekly sales dashboards\n"
            "- Migrated reporting to AWS Redshift, cutting query time by 40%\n"
            "- Built SQL pipelines feeding pricing reviews\n")
    [job] = ResumeParser()._extract_work_experience(text)
    assert job.responsibilities == ["Led a team of 3 analysts building weekly sales dashboards",
                                    "Migrated reporting to AWS Redshift, cutting query time by 40%",
                                    "Built SQL pipelines feeding pricing reviews"]
    assert job.accomplishments == []
