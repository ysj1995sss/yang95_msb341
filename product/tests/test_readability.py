"""Spec 010 readability check, on real generated PDFs (reportlab) read back with the same text
extraction the validator uses. Anonymized content."""

import io
import shutil

import pytest
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from resume_tailorer.artifacts.models import FindingSeverity
from resume_tailorer.pdf.readability import (
    ExpectedContent, check_readability, expected_for_docx,
)
from resume_tailorer.pdf.validator import PDFValidator, findings_from_pdf_issues
from tests.fixtures.ats import profile

NAME, EMAIL = "Riley Park", "riley.park@example.com"
B1 = "Built SQL dashboards used by 40 regional managers to track weekly campaign results"
B2 = "Led on-time delivery of a 6-month store launch across 4 teams, with weekly risk reviews"
B3 = "Raised email campaign conversion 12% by redesigning audience segments"
EXPECTED = ExpectedContent(name=NAME, email=EMAIL, bullets=(B1, B2, B3), employers=("Acme Retail", "Bluebird Goods"))


def _pdf(path, lines, image_line=None):
    """One line per row; `image_line` is drawn as a picture of text (no text layer)."""
    c = canvas.Canvas(str(path), pagesize=letter)
    y = 740
    for line in lines:
        if line == image_line:
            from PIL import Image, ImageDraw

            img = Image.new("RGB", (900, 24), "white")
            ImageDraw.Draw(img).text((2, 4), line, fill="black")
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            buf.seek(0)
            c.drawImage(ImageReader(buf), 54, y - 4, width=450, height=12)
        else:
            c.drawString(54, y, line)
        y -= 16
    c.save()
    return path


def _read(path):
    result = PDFValidator().validate(str(path), target_length="preserve")
    return result, check_readability(result.extracted_text, EXPECTED)


def _codes(findings, severity=None):
    return {f.code for f in findings if severity is None or f.severity == severity}


CLEAN = [NAME, EMAIL, "EXPERIENCE", "Marketing Analyst, Acme Retail, Jan 2022 - Present", f"- {B1}", f"- {B2}",
         f"- {B3}", "Marketing Coordinator, Bluebird Goods, Jun 2019 - Dec 2021", "EDUCATION",
         "BS Economics, State University, 2019"]


def test_a_clean_single_column_pdf_passes(tmp_path):
    _result, findings = _read(_pdf(tmp_path / "clean.pdf", CLEAN))
    assert findings == []


def test_a_bullet_drawn_as_an_image_fails_and_is_named(tmp_path):
    _result, findings = _read(_pdf(tmp_path / "image.pdf", CLEAN, image_line=f"- {B2}"))
    missing = [f for f in findings if f.code == "READABILITY_BULLET_MISSING"]
    assert missing and missing[0].severity == FindingSeverity.FAIL
    assert missing[0].details["bullets"] == [B2] and "Led on-time delivery" in missing[0].message


def test_two_column_text_that_reads_back_interleaved_is_warned_about(tmp_path):
    # A two-column page as many layouts write it: each row's left text, then the right column's.
    left = ["Built SQL dashboards used by 40 regional", "managers to track weekly campaign results",
            "Led on-time delivery of a 6-month store", "launch across 4 teams, with weekly risk reviews"]
    right = ["SKILLS", "SQL, Snowflake, Excel", "TOOLS", "Google Analytics"]
    rows = [f"{a}      {b}" for a, b in zip(left, right)]
    lines = [NAME, EMAIL, "EXPERIENCE", "Marketing Analyst, Acme Retail", *rows, f"- {B3}",
             "Bluebird Goods", "EDUCATION", "State University"]
    _result, findings = _read(_pdf(tmp_path / "columns.pdf", lines))
    order = [f for f in findings if f.code == "READABILITY_ORDER"]
    assert order and order[0].severity == FindingSeverity.WARNING
    assert "READABILITY_BULLET_MISSING" not in _codes(findings)  # the words are there, just out of order


def test_missing_contact_details_fail(tmp_path):
    _result, findings = _read(_pdf(tmp_path / "nocontact.pdf", CLEAN[2:]))
    contact = [f for f in findings if f.code == "READABILITY_CONTACT_MISSING"]
    assert contact and contact[0].severity == FindingSeverity.FAIL and contact[0].details["missing"] == ["name", "email"]


def test_a_blank_pdf_still_fails_text_not_extractable_unchanged(tmp_path):
    path = tmp_path / "blank.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.showPage()
    c.save()
    result, findings = _read(path)
    assert findings == []  # left to the existing check
    codes = {f.code: f for f in findings_from_pdf_issues(result.issues)}
    assert codes["TEXT_NOT_EXTRACTABLE"].severity == FindingSeverity.FAIL


def test_roles_out_of_page_order_unknown_symbols_and_unusual_headings_warn():
    text = (f"{NAME} {EMAIL}\nMy Journey\nBluebird Goods\n- {B1}\n- {B2}\n- {B3}\nAcme Retail �\n"
            "State University")
    findings = check_readability(text, EXPECTED)
    warnings = _codes(findings, FindingSeverity.WARNING)
    assert {"READABILITY_ROLE_ORDER", "READABILITY_UNMAPPED_SYMBOLS", "READABILITY_HEADING_UNUSUAL"} <= warnings
    assert not _codes(findings, FindingSeverity.FAIL)


def test_extraction_noise_is_tolerated():
    text = (f"{NAME}\n{EMAIL}\nEXPERIENCE\n- Built SQL dash-\nboards used by 40 regional managers to track weekly "
            f"campaign results.\n- {B2}\n- Raised email campaign conversion 12% by redesigning audience segﬁments\n"
            "EDUCATION")
    findings = check_readability(text.replace("segﬁments", "segments"), EXPECTED)
    assert not _codes(findings, FindingSeverity.FAIL)


def test_the_screen_lists_only_checks_that_ran_with_their_state():
    from resume_tailorer.ui.requirement_review_view import readability_view

    findings = check_readability("Riley Park\nMy Journey", EXPECTED)
    docx = readability_view(findings, "DOCX", ("pdf_technical", "readability"))
    labels = {i["label"]: i["state"] for i in docx["items"]}
    assert labels["Every bullet reads back"] == "fail" and labels["Standard section headings"] == "warn"
    assert "not a check by any employer's ATS" in docx["note"]
    freeform = {i["label"] for i in readability_view(findings, "PDF", ("readability",))["items"]}
    assert "Every bullet reads back" in freeform  # spec 011: the free-form path checks bullets too
    assert not any(label.startswith("Word file") for label in freeform)  # no Word file there
    old = readability_view([], "DOCX", ("pdf_technical",))  # a run from before spec 010
    assert {i["state"] for i in old["items"]} >= {"not_checked"}


def test_word_path_runs_the_readability_check_on_its_pdf(tmp_path, monkeypatch):
    """The Word pipeline's PDF missing a bullet now fails validation (it only checked that the
    PDF had text before)."""
    from docx import Document

    from resume_tailorer.analyzers.gap_analyzer import GapReport
    from resume_tailorer.analyzers.job_analyzer import JobAnalyzer
    from resume_tailorer.docx_export import pipeline as pipeline_module
    from resume_tailorer.docx_export.pipeline import run_docx_tailoring_pipeline

    doc = Document()
    for line in [NAME, EMAIL, "EXPERIENCE", "Marketing Analyst", "Acme Retail | Denver, CO | Jan 2022 - Present"]:
        doc.add_paragraph(line)
    for b in (B1, B2, B3):
        doc.add_paragraph(f"• {b}")
    for line in ["EDUCATION", "BS Economics, State University, 2019"]:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)

    pdf_missing_b2 = _pdf(tmp_path / "out.pdf", [line for line in CLEAN if B2 not in line])
    monkeypatch.setattr(pipeline_module, "convert_docx_to_pdf", lambda src, dst: shutil.copy(pdf_missing_b2, dst))

    class Keep:
        def tailor_bullets(self, bullets, *args, **kwargs):
            from resume_tailorer.tailorer.docx_bullet_tailorer import BulletEdit, BulletTailoringResult

            return BulletTailoringResult(edits=[BulletEdit(b.paragraph_index, b.text, b.text, False) for b in bullets],
                                         warnings=[])

    result = run_docx_tailoring_pipeline(buf.getvalue(), profile(), JobAnalyzer().analyze("Analyst"),
                                         GapReport([], ""), bullet_tailorer=Keep())
    codes = {f.code: f for f in result.validation.findings}
    assert codes["READABILITY_BULLET_MISSING"].severity == FindingSeverity.FAIL
    assert result.validation.status.value == "FAIL"
    assert "readability" in result.validation.checks_run


def test_expected_for_docx_skips_short_lines():
    expected = expected_for_docx(profile(), ["Skills", B1, ""])
    assert expected.bullets == (B1,) and expected.name == "Riley Park"


def test_a_second_reader_that_disagrees_is_warned_about_and_a_bullet_needs_both_to_miss_it():
    one_reader = f"{NAME} {EMAIL}\nEXPERIENCE\n- {B1}\n- {B2}\n- {B3}\nEDUCATION"
    scrambled = (f"{NAME} {EMAIL}\nEXPERIENCE\n- Built SQL dashboards SKILLS used by 40 TOOLS regional managers to\n"
                 f"- {B2}\n- {B3}")
    findings = check_readability(one_reader, EXPECTED, second_text=scrambled)
    disagree = [f for f in findings if f.code == "READABILITY_READERS_DISAGREE"]
    assert disagree and disagree[0].severity == FindingSeverity.WARNING and disagree[0].details["bullets"] == [B1]
    # Missing for one reader, present for the other: not a missing bullet.
    without_b3 = one_reader.replace(f"- {B3}", "")
    assert "READABILITY_BULLET_MISSING" not in _codes(check_readability(without_b3, EXPECTED, second_text=one_reader))


def test_the_second_reader_reads_a_real_pdf(tmp_path):
    from resume_tailorer.pdf.readability import second_reader_text

    text = second_reader_text(str(_pdf(tmp_path / "clean.pdf", CLEAN)))
    assert text and B1 in " ".join(text.split())
    assert second_reader_text(str(tmp_path / "missing.pdf")) is None


def test_word_file_checks_flag_header_contact_tables_and_text_boxes():
    from docx import Document
    from docx.oxml import parse_xml

    from resume_tailorer.pdf.readability import check_docx_file

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = f"{NAME} | {EMAIL}"
    doc.add_paragraph("EXPERIENCE")
    table = doc.add_table(rows=2, cols=2)
    for cell in [c for row in table.rows for c in row.cells]:
        cell.text = B1 + " " + B2
    doc.add_paragraph("Notes")._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><w:t>x</w:t>'
        '<wps:txbx><w:txbxContent><w:p><w:r><w:t>Skills in a box</w:t></w:r></w:p></w:txbxContent></wps:txbx></w:r>'))
    buf = io.BytesIO()
    doc.save(buf)
    codes = {f.code for f in check_docx_file(buf.getvalue(), EMAIL, NAME)}
    assert {"DOCX_CONTACT_IN_HEADER", "DOCX_TEXT_IN_TABLES", "DOCX_TEXT_BOXES"} <= codes

    clean = Document()
    for line in (NAME, EMAIL, "EXPERIENCE", f"- {B1}"):
        clean.add_paragraph(line)
    buf = io.BytesIO()
    clean.save(buf)
    assert check_docx_file(buf.getvalue(), EMAIL, NAME) == []


def test_free_form_pdfs_get_per_bullet_checks(tmp_path):
    from resume_tailorer.pdf.readability import expected_for_freeform

    tailored = "\n".join(CLEAN)
    expected = expected_for_freeform(profile(), tailored)
    assert expected.bullets == (B1, B2, B3) and not expected.check_contact
    pdf = _pdf(tmp_path / "freeform.pdf", [line for line in CLEAN if B2 not in line])
    result = PDFValidator().validate_artifact(str(pdf), profile=profile(), expected_page_count=None,
                                              accepted_changes=[], tailored_text=tailored)
    assert any(f.code == "READABILITY_BULLET_MISSING" for f in result.findings)
