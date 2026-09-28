from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from resume_tailorer.docx_export.layout_signature import (
    capture_layout_signature,
    compare_layout_signatures,
)
from resume_tailorer.docx_export.splicer import splice_bullets_into_docx
from resume_tailorer.tailorer.docx_bullet_tailorer import BulletEdit


def _document_with_target() -> Document:
    doc = Document()
    doc.add_paragraph("WORK EXPERIENCE", style="Heading 1")
    paragraph = doc.add_paragraph("Old bullet", style="List Paragraph")
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(num_id)
    p_pr.append(num_pr)
    doc.add_paragraph("EDUCATION", style="Heading 1")
    return doc


def test_text_replacement_preserves_immutable_layout_signature():
    doc = _document_with_target()
    before = capture_layout_signature(doc, editable_paragraph_indices={1})
    splice_bullets_into_docx(doc, [BulletEdit(1, "Old bullet", "New bullet", True)])
    after = capture_layout_signature(doc, editable_paragraph_indices={1})

    assert compare_layout_signatures(before, after) == []


def test_paragraph_reorder_is_a_failure():
    doc = _document_with_target()
    before = capture_layout_signature(doc, editable_paragraph_indices={1})
    body = doc._body._body
    body.insert(0, doc.paragraphs[-1]._p)

    findings = compare_layout_signatures(
        before, capture_layout_signature(doc, editable_paragraph_indices={1})
    )

    assert any(finding.code == "DOCX_PARAGRAPH_ORDER_CHANGED" for finding in findings)


def test_margin_change_is_a_failure():
    doc = _document_with_target()
    before = capture_layout_signature(doc, editable_paragraph_indices={1})
    doc.sections[0].left_margin = doc.sections[0].left_margin + 1000

    findings = compare_layout_signatures(
        before, capture_layout_signature(doc, editable_paragraph_indices={1})
    )

    assert any(finding.code == "DOCX_SECTION_LAYOUT_CHANGED" for finding in findings)
