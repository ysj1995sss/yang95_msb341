"""The shared tailoring service (spec 009): the same pipeline for Streamlit and the API.
Synthetic resume; the model is a stub, so nothing leaves the machine."""

import io
import json
import re

import pytest
from docx import Document

from resume_tailorer.artifacts.models import ChangeDisposition
from resume_tailorer.tailoring_service import TailoringError, regenerate, run_tailoring

JD = "Data Analyst. Requirements: SQL, Tableau and dashboard reporting for business partners."


def _resume() -> bytes:
    doc = Document()
    for line in ("Riley Park", "riley@example.com | 555-0100", "", "EXPERIENCE", "Data Analyst",
                 "Acme Analytics | Denver, CO | 2021 - Present"):
        doc.add_paragraph(line)
    doc.add_paragraph("• Built SQL dashboards used by 40 managers across regional sales teams")
    doc.add_paragraph("• Cut weekly reporting time by 30%")
    for line in ("", "EDUCATION", "BS Economics, State University, 2019", "", "SKILLS", "SQL, Tableau"):
        doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


class RewritingModel:
    """Rewrites the first bullet it sees, keeps the rest (the bullet tailorer's JSON shape)."""

    def complete(self, system, user, max_tokens=2000):
        found = re.search(r'\[\s*\{\s*"paragraph_index"', user)
        if not found:
            return "[]"  # other passes (length correction): no further edits
        bullets, _end = json.JSONDecoder().raw_decode(user[found.start():])
        edits = [{"paragraph_index": b["paragraph_index"], "change": "keep", "new_text": ""} for b in bullets]
        for edit, bullet in zip(edits, bullets):
            if bullet["text"].startswith("Built SQL dashboards"):
                edit.update(change="rewrite", new_text="Built Tableau SQL dashboards used by 40 managers across regional sales teams")
        return json.dumps(edits)


@pytest.fixture
def session(tmp_path):
    return {"artifacts_dir": str(tmp_path)}


def test_a_word_resume_is_tailored_and_rebuilt_from_decisions(session, monkeypatch):
    from resume_tailorer.docx_export import converter

    # No office suite here: the pipeline must still produce a reviewable result.
    monkeypatch.setattr(converter, "_word_platform", lambda: False)
    monkeypatch.setattr(converter, "_soffice", lambda: None)
    steps = []
    state = run_tailoring(session, original_bytes=_resume(), filename="riley.docx", job_description=JD,
                          pending={"job_id": "greenhouse_1", "title": "Data Analyst", "company": "Acme"},
                          llm=RewritingModel(), progress=steps.append)
    assert steps[0] == "Reading your resume" and len(steps) == 3
    assert state["source_kind"] == "DOCX" and state["version"] == 1
    proposed = [c for c in state["changes"] if "Tableau" in (c.proposed_text or "")]
    assert proposed, [c.proposed_text for c in state["changes"]]

    state["dispositions"][proposed[0].change_id] = ChangeDisposition.REJECTED.value
    regenerate(session, state)
    assert state["version"] == 2 and state["reviewed"] is True and state["dirty"] is False
    assert "Built SQL dashboards used by 40 managers across regional sales teams" in state["tailored_text"]


def test_an_empty_manual_edit_is_refused(session, monkeypatch):
    from resume_tailorer.docx_export import converter

    monkeypatch.setattr(converter, "_word_platform", lambda: False)
    monkeypatch.setattr(converter, "_soffice", lambda: None)
    state = run_tailoring(session, original_bytes=_resume(), filename="riley.docx", job_description=JD,
                          pending={}, llm=RewritingModel())
    change = next(c for c in state["changes"] if "Tableau" in (c.proposed_text or ""))
    state["dispositions"][change.change_id] = ChangeDisposition.MANUALLY_EDITED.value
    state["manual_texts"][change.change_id] = "   "
    with pytest.raises(TailoringError, match="empty"):
        regenerate(session, state)


def test_no_job_description_is_a_plain_message(session):
    with pytest.raises(TailoringError, match="job description"):
        run_tailoring(session, original_bytes=b"x", filename="r.docx", job_description=" ", pending={},
                      llm=RewritingModel())
