"""A real Word-to-PDF conversion with LibreOffice, as on the Linux server (decision 027).
Runs where LibreOffice is installed (CI installs packages.txt); skipped elsewhere."""

import pytest
from docx import Document
from pypdf import PdfReader

from resume_tailorer.docx_export import converter

pytestmark = pytest.mark.skipif(converter._soffice() is None, reason="LibreOffice is not installed here")


def test_a_word_resume_becomes_a_one_page_pdf_with_its_text(tmp_path, monkeypatch):
    monkeypatch.setattr(converter, "_word_platform", lambda: False)  # take the server's path
    doc = Document()
    doc.add_paragraph("Riley Park")
    doc.add_paragraph("EXPERIENCE")
    doc.add_paragraph("• Built SQL dashboards used by 40 managers")
    source, target = tmp_path / "resume.docx", tmp_path / "resume.pdf"
    doc.save(source)
    converter.convert_docx_to_pdf(str(source), str(target))
    reader = PdfReader(str(target))
    assert len(reader.pages) == 1
    assert "Built SQL dashboards" in reader.pages[0].extract_text()
