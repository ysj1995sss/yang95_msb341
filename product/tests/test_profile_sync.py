"""Spec 011: Career Profile edits reach a Word resume before tailoring, keeping its layout.
Anonymized, synthetic documents."""

import io

from docx import Document

from resume_tailorer.docx_export.profile_sync import sync_docx_to_profile
from resume_tailorer.models import CareerTruthProfile, EducationEntry, WorkExperience
from resume_tailorer.parsers.docx_structure import is_bullet_paragraph

A1 = "Built SQL dashboards used by 40 regional managers"
A2 = "Led on-time delivery of a 6-month store launch across 4 teams"
A3 = "Ran weekly campaign reviews with merchandising"
B1 = "Cleaned Salesforce campaign data before each quarterly review"


def _docx(typed=False) -> bytes:
    doc = Document()
    for line in ("Riley Park", "riley@example.com", "EXPERIENCE", "Marketing Analyst", "Acme Retail | Jan 2022 - Present"):
        doc.add_paragraph(line)
    for b in (A1, A2, A3):
        doc.add_paragraph(f"• {b}") if typed else doc.add_paragraph(b, style="List Bullet")
    for line in ("Marketing Coordinator", "Bluebird Goods | Jun 2019 - Dec 2021"):
        doc.add_paragraph(line)
    doc.add_paragraph(f"• {B1}") if typed else doc.add_paragraph(B1, style="List Bullet")
    for line in ("EDUCATION", "BS Economics, State University, 2019"):
        doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _profile(acme_bullets, extra_role=False) -> CareerTruthProfile:
    jobs = [  # deliberately NOT in the file's order
        WorkExperience(employer="Bluebird Goods", title="Marketing Coordinator", dates="Jun 2019 - Dec 2021",
                       responsibilities=[B1], accomplishments=[]),
        WorkExperience(employer="Acme Retail", title="Marketing Analyst", dates="Jan 2022 - Present",
                       responsibilities=list(acme_bullets), accomplishments=[]),
    ]
    if extra_role:
        jobs.append(WorkExperience(employer="Northwind Outdoor", title="Strategy Intern", dates="2024",
                                   responsibilities=["Sized a new market"], accomplishments=[]))
    return CareerTruthProfile(contact_info={"name": "Riley Park"}, education=[EducationEntry("BS", "Economics", "State University", 2019)],
                              work_experience=jobs, skills=[], tools=[], certifications=[], accomplishments=[])


def _texts(docx_bytes):
    return [p.text for p in Document(io.BytesIO(docx_bytes)).paragraphs]


def _two_acme_roles() -> bytes:
    doc = Document(io.BytesIO(_docx()))
    education = next(p for p in doc.paragraphs if p.text == "EDUCATION")
    education.insert_paragraph_before("Sales Analyst")
    education.insert_paragraph_before("Acme Retail | Jan 2020 - Dec 2021")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_unique_employer_and_unchanged_dates_identify_a_renamed_title():
    profile = _profile([A1, A2, A3])
    profile.work_experience[1].title = "Senior Marketing Analyst"
    result = sync_docx_to_profile(_docx(), profile)
    texts = _texts(result.docx_bytes)
    assert "Senior Marketing Analyst" in texts
    assert "Marketing Analyst" not in texts
    assert any(f.field == "work.title" and f.outcome == "updated" for f in result.fields)


def test_two_roles_at_one_employer_do_not_guess_a_renamed_title():
    profile = _profile([A1, A2, A3])
    profile.work_experience[1].title = "Senior Marketing Analyst"
    profile.work_experience.append(WorkExperience(
        employer="Acme Retail", title="Sales Analyst", dates="Jan 2020 - Dec 2021",
        responsibilities=[], accomplishments=[],
    ))
    result = sync_docx_to_profile(_two_acme_roles(), profile)
    assert result.has_blockers
    assert any(f.field == "work.title" and "ambiguous" in f.reason.lower() for f in result.fields)
    assert "Senior Marketing Analyst" not in _texts(result.docx_bytes)


def test_summary_skills_and_tools_sync_to_their_existing_sections():
    doc = Document(io.BytesIO(_docx()))
    experience = next(p for p in doc.paragraphs if p.text == "EXPERIENCE")
    experience.insert_paragraph_before("Old summary about retail reporting")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Skills: SQL")
    doc.add_paragraph("Tools: Excel")
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.summary = "Marketing analyst focused on retail reporting"
    profile.skills = ["SQL", "Tableau"]
    profile.tools = ["Excel", "Power BI"]
    result = sync_docx_to_profile(out.getvalue(), profile)
    texts = _texts(result.docx_bytes)
    assert profile.summary in texts
    assert "Skills: SQL, Tableau" in texts
    assert "Tools: Excel, Power BI" in texts
    assert texts.index("EXPERIENCE") < texts.index("EDUCATION") < texts.index("SKILLS")


def test_missing_skills_section_reports_unplaced_fact_without_inserting_it():
    profile = _profile([A1, A2, A3])
    profile.skills = ["Tableau"]
    result = sync_docx_to_profile(_docx(), profile)
    assert any(f.field == "skills" and f.outcome == "unplaced" for f in result.fields)
    assert "Tableau" not in " ".join(_texts(result.docx_bytes))


def test_unique_header_email_syncs_in_place():
    doc = Document(io.BytesIO(_docx()))
    doc.paragraphs[1].text = ""
    doc.sections[0].header.paragraphs[0].text = "riley.old@example.com"
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.contact_info["email"] = "riley.new@example.com"
    result = sync_docx_to_profile(out.getvalue(), profile)
    changed = Document(io.BytesIO(result.docx_bytes))
    assert changed.sections[0].header.paragraphs[0].text == "riley.new@example.com"
    assert any(f.field == "contact.email" and f.outcome == "updated" for f in result.fields)


def test_education_degree_and_year_sync_without_losing_field():
    profile = _profile([A1, A2, A3])
    profile.education[0].degree = "Bachelor of Science"
    profile.education[0].year = 2020
    result = sync_docx_to_profile(_docx(), profile)
    assert "Bachelor of Science Economics, State University, 2020" in _texts(result.docx_bytes)


def test_duplicate_institution_does_not_swap_education_entries():
    doc = Document(io.BytesIO(_docx()))
    doc.add_paragraph("MS Marketing, State University, 2021")
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.education = [
        EducationEntry("Bachelor of Science", "Economics", "State University", 2020),
        EducationEntry("Master of Science", "Marketing", "State University", 2022),
    ]
    result = sync_docx_to_profile(out.getvalue(), profile)
    assert result.has_blockers
    assert "BS Economics, State University, 2019" in _texts(result.docx_bytes)
    assert "MS Marketing, State University, 2021" in _texts(result.docx_bytes)


def test_missing_certifications_section_is_reported_without_insertion():
    profile = _profile([A1, A2, A3])
    profile.certifications = ["PMP"]
    result = sync_docx_to_profile(_docx(), profile)
    assert any(f.field == "certifications" and f.outcome == "unplaced" for f in result.fields)
    assert "PMP" not in " ".join(_texts(result.docx_bytes))


def test_two_education_entries_at_one_school_match_by_degree():
    doc = Document(io.BytesIO(_docx()))
    doc.add_paragraph("MS Marketing, State University, 2021")
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.education.append(EducationEntry("MS", "Marketing", "State University", 2022))
    result = sync_docx_to_profile(out.getvalue(), profile)
    assert not result.has_blockers
    assert "BS Economics, State University, 2019" in _texts(result.docx_bytes)
    assert "MS Marketing, State University, 2022" in _texts(result.docx_bytes)


def test_job_dates_change_without_discarding_existing_location():
    doc = Document(io.BytesIO(_docx()))
    next(p for p in doc.paragraphs if p.text == "Acme Retail | Jan 2022 - Present").text = (
        "Acme Retail | Denver, CO | Jan 2022 - Present"
    )
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.work_experience[1].dates = "Jan 2023 - Present"
    profile.work_experience[1].location = "Denver, CO"
    result = sync_docx_to_profile(out.getvalue(), profile)
    assert "Acme Retail | Denver, CO | Jan 2023 - Present" in _texts(result.docx_bytes)


def test_changed_employer_with_same_title_is_reported_as_blocking():
    profile = _profile([A1, A2, A3])
    profile.work_experience[1].employer = "Cedar Retail"
    result = sync_docx_to_profile(_docx(), profile)
    assert result.has_blockers
    assert any(f.field == "work.employer" for f in result.fields)


def test_single_skills_line_under_skills_heading_is_updated_in_place():
    doc = Document(io.BytesIO(_docx()))
    doc.add_paragraph("SKILLS")
    old = doc.add_paragraph("SQL, Excel")
    old.style = "List Paragraph"
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.skills = ["SQL", "Tableau"]
    result = sync_docx_to_profile(out.getvalue(), profile)
    changed = Document(io.BytesIO(result.docx_bytes))
    line = next(p for p in changed.paragraphs if p.text == "SQL, Tableau")
    assert line.style.name == "List Paragraph"


def test_contact_location_line_is_updated_without_moving_it():
    doc = Document(io.BytesIO(_docx()))
    doc.paragraphs[2].insert_paragraph_before("Denver, CO")
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.contact_info["location"] = "Boulder, CO"
    result = sync_docx_to_profile(out.getvalue(), profile)
    texts = _texts(result.docx_bytes)
    assert "Boulder, CO" in texts
    assert "Denver, CO" not in texts
    assert any(f.field == "contact.location" and f.outcome == "updated" for f in result.fields)


def test_location_in_combined_contact_line_is_replaced_without_losing_email_or_phone():
    doc = Document(io.BytesIO(_docx()))
    doc.paragraphs[1].text = "riley@example.com | 555-0100 | Denver, CO"
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.contact_info["location"] = "Boulder, CO"
    result = sync_docx_to_profile(out.getvalue(), profile)
    assert "riley@example.com | 555-0100 | Boulder, CO" in _texts(result.docx_bytes)
    assert not result.has_blockers


def test_existing_certifications_line_is_updated_not_moved():
    doc = Document(io.BytesIO(_docx()))
    doc.add_paragraph("CERTIFICATIONS")
    doc.add_paragraph("Old Certificate")
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.certifications = ["PMP"]
    result = sync_docx_to_profile(out.getvalue(), profile)
    assert not result.has_blockers
    assert _texts(result.docx_bytes)[-2:] == ["CERTIFICATIONS", "PMP"]


def test_core_competencies_heading_is_a_safe_skills_target():
    doc = Document(io.BytesIO(_docx()))
    doc.add_paragraph("CORE COMPETENCIES")
    doc.add_paragraph("SQL, Excel")
    out = io.BytesIO()
    doc.save(out)
    profile = _profile([A1, A2, A3])
    profile.skills = ["SQL", "Tableau"]
    result = sync_docx_to_profile(out.getvalue(), profile)
    assert "SQL, Tableau" in _texts(result.docx_bytes)


def test_ambiguous_stale_role_blocks_tailoring_before_the_model(tmp_path):
    import pytest

    from resume_tailorer.tailoring_service import TailoringError, run_tailoring

    profile = _profile([A1, A2, A3])
    profile.work_experience[1].title = "Senior Marketing Analyst"
    profile.work_experience.append(WorkExperience(
        employer="Acme Retail", title="Sales Analyst", dates="Jan 2020 - Dec 2021",
        responsibilities=[], accomplishments=[],
    ))

    class FailIfCalled:
        def complete(self, system, user, max_tokens=2000):
            raise AssertionError("Blocked Word runs must not call the model")

    session = {"artifacts_dir": str(tmp_path)}
    with pytest.raises(TailoringError, match="Word file"):
        run_tailoring(session, original_bytes=_two_acme_roles(), filename="riley.docx",
                      job_description="Marketing analyst with SQL reporting experience",
                      pending={"job_id": "j1", "title": "Analyst", "company": "Acme"},
                      llm=FailIfCalled(), career_profile=profile)
    assert not session.get("active_handoff")


def test_an_unchanged_profile_leaves_the_file_byte_for_byte():
    original = _docx()
    result = sync_docx_to_profile(original, _profile([A1, A2, A3]))
    assert result.docx_bytes == original and result.summary == ""


def test_edits_additions_and_removals_reach_the_file_in_its_own_order():
    edited = A2.replace("store launch", "store launch and Tableau reporting")
    new = "Built Tableau dashboards for the loyalty program"
    result = sync_docx_to_profile(_docx(), _profile([new, edited, A1]))  # A3 removed, order differs
    texts = _texts(result.docx_bytes)
    acme = texts[texts.index("Acme Retail | Jan 2022 - Present") + 1: texts.index("Marketing Coordinator")]
    assert acme == [A1, edited, new]  # the file's order kept; the new bullet after the last one
    assert (result.changed, result.added, result.removed) == (1, 1, 1)
    assert B1 in texts  # the other role untouched
    assert "1 bullet(s) updated, 1 added, 1 removed" in result.summary


def test_an_added_bullet_keeps_the_roles_bullet_style_and_numbering():
    result = sync_docx_to_profile(_docx(), _profile([A1, A2, A3, "Built Tableau dashboards for the loyalty program"]))
    doc = Document(io.BytesIO(result.docx_bytes))
    added = next(p for p in doc.paragraphs if p.text.startswith("Built Tableau"))
    assert added.style.name == "List Bullet" and is_bullet_paragraph(added)


def test_typed_bullets_keep_their_marker():
    result = sync_docx_to_profile(_docx(typed=True), _profile([A1, A2 + " on budget", A3]))
    assert f"• {A2} on budget" in _texts(result.docx_bytes)


def test_roles_are_matched_by_employer_not_position_and_the_profile_is_aligned_to_the_file():
    result = sync_docx_to_profile(_docx(), _profile([A1, A2, A3]))
    assert [j.employer for j in result.aligned_profile.work_experience] == ["Acme Retail", "Bluebird Goods"]


def test_roles_only_in_the_profile_or_only_in_the_file_are_reported():
    result = sync_docx_to_profile(_docx(), _profile([A1, A2, A3], extra_role=True))
    assert result.unmatched_profile_roles == ["Strategy Intern at Northwind Outdoor"]
    assert "Not in your Word file" in result.summary
    one_role = _profile([A1, A2, A3])
    one_role.work_experience = one_role.work_experience[1:]  # Bluebird removed from the profile
    result = sync_docx_to_profile(_docx(), one_role)
    assert result.unmatched_file_roles and "Bluebird Goods" in result.unmatched_file_roles[0]
    assert B1 in _texts(result.docx_bytes)  # left as it is


def test_a_profile_edit_reaches_the_tailored_resume_and_closes_the_dead_end(monkeypatch, tmp_path):
    """Adding 'Built Tableau dashboards...' to a role in Career Profile makes Tableau direct
    evidence on the next run, and the bullet is in the tailored Word file (it used to be ignored)."""
    from resume_tailorer.docx_export import converter
    from resume_tailorer.tailoring_service import run_tailoring

    monkeypatch.setattr(converter, "_word_platform", lambda: False)
    monkeypatch.setattr(converter, "_soffice", lambda: None)

    class KeepEverything:
        def complete(self, system, user, max_tokens=2000):
            return "[]"

    posting = ("Data Analyst at Northwind.\nRequirements:\n- SQL and Tableau dashboards for business partners.\n")
    new_bullet = "Built Tableau dashboards for the loyalty program"
    profile = _profile([A1, A2, A3, new_bullet])
    session = {"artifacts_dir": str(tmp_path)}
    state = run_tailoring(session, original_bytes=_docx(), filename="riley.docx", job_description=posting,
                          pending={"job_id": "j1", "title": "Data Analyst", "company": "Northwind"},
                          llm=KeepEverything(), career_profile=profile)
    row = next(r for r in state["requirement_review"].rows if r.text.startswith("SQL and Tableau"))
    assert dict(row.term_status)["Tableau"] == "direct"
    assert new_bullet in _texts(state["docx_bytes"])
    assert "1 added" in state["sync_summary"]
    # The session keeps the real Career Profile (not the file-ordered working copy).
    from resume_tailorer.session_profile import get_career_profile

    assert [j.employer for j in get_career_profile(session).work_experience] == ["Bluebird Goods", "Acme Retail"]
