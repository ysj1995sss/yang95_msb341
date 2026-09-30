import shutil

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from resume_tailorer.models import CareerTruthProfile, WorkExperience
from resume_tailorer.analyzers import JobAnalyzer
from resume_tailorer.analyzers.gap_analyzer import EvidenceLevel, GapCategory, GapItem, GapReport
from resume_tailorer.tailorer.docx_bullet_tailorer import BulletEdit, BulletTailoringResult
from resume_tailorer.docx_export.converter import DocxConversionUnavailable
from resume_tailorer.docx_export import pipeline as pipeline_module
from resume_tailorer.docx_export.pipeline import run_docx_tailoring_pipeline
from resume_tailorer.docx_export.pipeline import changes_from_bullet_edits
from resume_tailorer.docx_export.pipeline import finalize_docx_edits
from resume_tailorer.parsers.docx_structure import extract_docx_structure
from resume_tailorer.parsers.docx_structure import Bullet
from resume_tailorer.artifacts.models import ValidationStatus


def _add_bullet_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph(text, style="List Paragraph")
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    return paragraph


def _sample_docx_bytes(tmp_path) -> bytes:
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("A summary paragraph.")
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    doc.add_paragraph("Marketing Manager")
    doc.add_paragraph("Acme Corp | Remote\t2022-2024")
    _add_bullet_paragraph(doc, "Did a thing")
    path = tmp_path / "sample.docx"
    doc.save(path)
    return path.read_bytes()


@pytest.fixture
def sample_profile():
    return CareerTruthProfile(
        contact_info={"name": "Jane Doe"},
        education=[],
        work_experience=[
            WorkExperience(
                employer="Acme Corp",
                title="Marketing Manager",
                dates="2022-2024",
                responsibilities=[],
                accomplishments=[],
            )
        ],
        skills=[],
        tools=[],
        certifications=[],
        accomplishments=[],
    )


@pytest.fixture
def sample_job_analysis():
    return JobAnalyzer().analyze("Marketing role requiring SQL.")


@pytest.fixture
def sample_gap_report():
    return GapReport(items=[], summary="no gaps")


class _StubBulletTailorer:
    """Returns every bullet unchanged except paragraph 5 ('Did a thing')."""

    def tailor_bullets(self, bullets, profile, job_analysis, gap_report, max_repair_attempts=1, priority_focus=None):
        edits = []
        for b in bullets:
            if b.text == "Did a thing":
                edits.append(BulletEdit(b.paragraph_index, b.text, "Did a thing using SQL", changed=True))
            else:
                edits.append(BulletEdit(b.paragraph_index, b.text, b.text, changed=False))
        return BulletTailoringResult(edits=edits, warnings=[])


def test_docx_bytes_always_produced_even_when_conversion_unavailable(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report
):
    def fake_convert(src, dst):
        raise DocxConversionUnavailable("no Word on this machine")

    monkeypatch.setattr(pipeline_module, "convert_docx_to_pdf", fake_convert)

    original_bytes = _sample_docx_bytes(tmp_path)
    result = run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )

    assert result.docx_bytes  # splicing must never depend on conversion succeeding
    assert result.conversion_available is False
    assert result.pdf_bytes is None
    assert result.original_page_count is None
    assert result.tailored_page_count is None
    assert result.page_count_preserved is None
    assert any("Word" in w or "unavailable" in w.lower() for w in result.bullet_warnings)


def test_splice_applied_to_docx_bytes(tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report):
    monkeypatch.setattr(
        pipeline_module, "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip conversion in this test")),
    )
    original_bytes = _sample_docx_bytes(tmp_path)
    result = run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )

    import io
    spliced_doc = Document(io.BytesIO(result.docx_bytes))
    bullet_texts = [p.text for p in spliced_doc.paragraphs if p.text == "Did a thing using SQL"]
    assert bullet_texts == ["Did a thing using SQL"]


def test_conversion_success_reports_page_counts(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report
):
    """Simulate a successful conversion by copying a tiny pre-built PDF into
    place, without touching real Word -- this is the mockable seam the
    pipeline is designed around."""
    fixture_pdf = tmp_path / "fixture.pdf"
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(fixture_pdf), pagesize=(200, 200))
    c.drawString(10, 100, "Hello World")
    c.showPage()
    c.save()

    def fake_convert(src, dst):
        shutil.copy(fixture_pdf, dst)

    monkeypatch.setattr(pipeline_module, "convert_docx_to_pdf", fake_convert)

    original_bytes = _sample_docx_bytes(tmp_path)
    result = run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )

    assert result.conversion_available is True
    assert result.original_page_count == 1
    assert result.tailored_page_count == 1
    assert result.page_count_preserved is True
    assert result.pdf_bytes is not None


def test_convert_to_pdf_false_skips_conversion_entirely(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report
):
    def fail_if_called(src, dst):
        raise AssertionError("convert_docx_to_pdf should not be called when convert_to_pdf=False")

    monkeypatch.setattr(pipeline_module, "convert_docx_to_pdf", fail_if_called)
    original_bytes = _sample_docx_bytes(tmp_path)
    result = run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
        convert_to_pdf=False,
    )
    assert result.conversion_available is False
    assert result.pdf_bytes is None
    assert result.docx_bytes


def test_docx_result_uses_bullet_edit_pairing_not_similarity_matching():
    bullet = Bullet(7, "Original exact bullet", "work_experience", job_index=0)
    edit = BulletEdit(7, "Original exact bullet", "Tailored exact bullet", True)

    changes = changes_from_bullet_edits(
        [edit], [bullet], GapReport(items=[], summary="No gaps")
    )

    assert changes[0].source_index == 7
    assert changes[0].original_text == "Original exact bullet"
    assert changes[0].proposed_text == "Tailored exact bullet"


def test_pipeline_exposes_structured_docx_validation(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report
):
    monkeypatch.setattr(
        pipeline_module,
        "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip")),
    )

    result = run_docx_tailoring_pipeline(
        _sample_docx_bytes(tmp_path),
        sample_profile,
        sample_job_analysis,
        sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )

    assert result.validation.status in {ValidationStatus.PASS, ValidationStatus.WARNING}
    assert result.changes[0].source_index == 5


def test_shallow_tailoring_is_flagged_when_evidence_goes_unused(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis
):
    """Problem 7's self-check: found live (2026-09-23) -- a real pass
    changed only 2 bullets while 11 requirements had real evidence
    elsewhere in the resume. This should surface a warning, not silently
    ship a shallow result."""
    monkeypatch.setattr(
        pipeline_module, "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip")),
    )
    gap_report = GapReport(
        items=[
            GapItem(requirement=f"Requirement {i}", category=GapCategory.B, reason="evidence exists", candidate_evidence="x")
            for i in range(5)
        ],
        summary="5 gaps found",
    )
    original_bytes = _sample_docx_bytes(tmp_path)
    result = run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )
    assert result.bullets_changed == 1
    assert result.addressable_requirements == 5
    assert result.tailoring_seems_shallow is True
    assert any("shallow" in w.lower() for w in result.bullet_warnings)


def test_not_flagged_shallow_when_little_evidence_exists(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report
):
    monkeypatch.setattr(
        pipeline_module, "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip")),
    )
    original_bytes = _sample_docx_bytes(tmp_path)
    result = run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )
    assert result.tailoring_seems_shallow is False


def test_job_count_mismatch_falls_back_without_crashing(
    tmp_path, monkeypatch, sample_job_analysis, sample_gap_report
):
    """If the profile has a different number of jobs than the DOCX
    structure walker found, scoring-text synthesis must degrade gracefully
    rather than index out of range or mispair a job's bullets."""
    monkeypatch.setattr(
        pipeline_module, "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip")),
    )
    mismatched_profile = CareerTruthProfile(
        contact_info={"name": "Jane Doe"},
        education=[],
        work_experience=[],  # zero jobs, but the docx has one
        skills=[],
        tools=[],
        certifications=[],
        accomplishments=[],
    )
    original_bytes = _sample_docx_bytes(tmp_path)
    result = run_docx_tailoring_pipeline(
        original_bytes, mismatched_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )
    assert "WORK EXPERIENCE" in result.tailored_scoring_text


class _SecondPassBulletTailorer:
    """First call: shallow, changes nothing. Second call (only reached if
    priority_focus is passed): makes a genuine additional change to 'Did
    a thing', proving the resume-wide optimization pass actually improves
    the result rather than just re-running the same prompt."""

    def __init__(self):
        self.calls = []

    def tailor_bullets(self, bullets, profile, job_analysis, gap_report, max_repair_attempts=1, priority_focus=None):
        self.calls.append(priority_focus)
        if priority_focus is None:
            edits = [BulletEdit(b.paragraph_index, b.text, b.text, changed=False) for b in bullets]
            return BulletTailoringResult(edits=edits, warnings=[])
        edits = []
        for b in bullets:
            if b.text == "Did a thing":
                edits.append(BulletEdit(b.paragraph_index, b.text, "Did a thing, using SQL for analysis", changed=True))
            else:
                edits.append(BulletEdit(b.paragraph_index, b.text, b.text, changed=False))
        return BulletTailoringResult(edits=edits, warnings=[])


def test_resume_wide_optimization_pass_runs_when_shallow_and_strong_evidence_remains(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis
):
    """Steps 10-15 audit Phase E: a shallow first pass with strong,
    unrepresented evidence should trigger a bounded second pass that
    surfaces it -- Step 14's resume-wide optimization question answered
    concretely, not left as a warning nobody acts on."""
    monkeypatch.setattr(
        pipeline_module, "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip")),
    )
    gap_report = GapReport(
        items=[
            GapItem(
                requirement=f"Requirement {i}", category=GapCategory.B,
                reason="strong evidence", candidate_evidence="x",
                evidence_level=EvidenceLevel.STRONGLY_SUPPORTED,
            )
            for i in range(5)
        ],
        summary="5 gaps found",
    )
    original_bytes = _sample_docx_bytes(tmp_path)
    tailorer = _SecondPassBulletTailorer()
    result = run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, gap_report,
        bullet_tailorer=tailorer,
    )

    assert len(tailorer.calls) == 2, "expected exactly one first pass and one optimization pass"
    assert tailorer.calls[0] is None
    assert tailorer.calls[1] is not None and len(tailorer.calls[1]) > 0
    assert result.bullets_changed == 1
    assert "Did a thing, using SQL for analysis" in result.tailored_scoring_text
    assert any("optimization pass" in w.lower() for w in result.bullet_warnings)


def test_resume_wide_optimization_pass_does_not_run_when_nothing_to_focus_on(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report
):
    """sample_gap_report has zero items -- shallow can't even be true, so
    the optimization pass must not fire a second LLM call at all."""
    monkeypatch.setattr(
        pipeline_module, "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip")),
    )
    original_bytes = _sample_docx_bytes(tmp_path)
    tailorer = _SecondPassBulletTailorer()
    run_docx_tailoring_pipeline(
        original_bytes, sample_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=tailorer,
    )
    assert len(tailorer.calls) == 1


def test_finalize_docx_edits_never_calls_the_tailorer(
    tmp_path, monkeypatch, sample_profile, sample_gap_report
):
    """Steps 16-20 regeneration reconstructs BulletEdits from stored,
    human-reviewed ResumeChange dispositions and splices them directly --
    finalize_docx_edits is the shared, tailorer-free half of the pipeline
    that makes this possible without re-invoking the LLM (which would
    silently discard the review)."""
    monkeypatch.setattr(
        pipeline_module, "convert_docx_to_pdf",
        lambda src, dst: (_ for _ in ()).throw(DocxConversionUnavailable("skip")),
    )
    import io

    original_bytes = _sample_docx_bytes(tmp_path)
    doc = Document(io.BytesIO(original_bytes))
    structure = extract_docx_structure(doc)
    bullets = structure.splice_targets(doc.paragraphs)
    target = next(b for b in bullets if b.text == "Did a thing")
    edits = [BulletEdit(target.paragraph_index, target.text, "Reverted to original", changed=True)]

    result = finalize_docx_edits(
        doc=doc,
        structure=structure,
        bullets=bullets,
        edits=edits,
        original_docx_bytes=original_bytes,
        gap_report=sample_gap_report,
        profile=sample_profile,
        convert_to_pdf=False,
    )

    spliced = Document(io.BytesIO(result.docx_bytes))
    assert any(p.text == "Reverted to original" for p in spliced.paragraphs)
    assert result.bullets_evaluated == 0  # caller-supplied, not re-derived


def test_blank_generated_pdf_fails_validation(
    tmp_path, monkeypatch, sample_profile, sample_job_analysis, sample_gap_report
):
    """A converted PDF with no extractable text must never be reported as PASS."""
    from reportlab.pdfgen import canvas

    blank_pdf = tmp_path / "blank.pdf"
    c = canvas.Canvas(str(blank_pdf), pagesize=(200, 200))
    c.showPage()
    c.save()
    monkeypatch.setattr(pipeline_module, "convert_docx_to_pdf", lambda src, dst: shutil.copy(blank_pdf, dst))

    result = run_docx_tailoring_pipeline(
        _sample_docx_bytes(tmp_path), sample_profile, sample_job_analysis, sample_gap_report,
        bullet_tailorer=_StubBulletTailorer(),
    )

    assert result.validation.status is ValidationStatus.FAIL
    assert "TEXT_NOT_EXTRACTABLE" in {f.code for f in result.validation.findings}
