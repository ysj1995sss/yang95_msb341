"""Conservative, section-aware Career Profile edits to an existing Word document."""

from __future__ import annotations

import re
from dataclasses import dataclass

from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

from resume_tailorer.models import CareerTruthProfile
from resume_tailorer.parsers.docx_structure import extract_docx_structure


@dataclass(frozen=True)
class SyncField:
    field: str
    value: str
    outcome: str
    reason: str = ""
    blocking: bool = False


_EMAIL = re.compile(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b")
_PHONE = re.compile(r"(?:\+?\d[\d(). -]{7,}\d)")
_LOCATION = re.compile(r"[A-Za-z][A-Za-z .'-]+,\s*[A-Z]{2}")
_SECTION = re.compile(r"^(?:SUMMARY|PROFESSIONAL SUMMARY|EXPERIENCE|WORK EXPERIENCE|EDUCATION|SKILLS|TECHNICAL SKILLS|CORE COMPETENCIES|TOOLS|CERTIFICATIONS|ADDITIONAL)$", re.I)


def rewrite_paragraph(paragraph: Paragraph, wanted: str) -> bool:
    """Keep the paragraph and its first run's style; leave an exact match untouched."""
    if paragraph.text == wanted:
        return False
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(wanted)
    else:
        runs[0].text = wanted
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    return True


def _all_paragraphs(doc: DocxDocument) -> list[Paragraph]:
    result = list(doc.paragraphs)
    for section in doc.sections:
        result.extend(section.header.paragraphs)
        result.extend(section.footer.paragraphs)
    return result


def _record_update(fields: list[SyncField], field: str, value: str, paragraph: Paragraph, wanted: str) -> None:
    if rewrite_paragraph(paragraph, wanted):
        fields.append(SyncField(field, value, "updated"))


def _unplaced(fields: list[SyncField], field: str, value: str, reason: str, *, blocking: bool = False) -> None:
    fields.append(SyncField(field, value, "unplaced", reason, blocking))


def _section_body(doc: DocxDocument, heading: str) -> list[Paragraph]:
    paragraphs = doc.paragraphs
    matches = [i for i, p in enumerate(paragraphs) if p.text.strip().lower() == heading.lower()]
    if len(matches) != 1:
        return []
    start = matches[0] + 1
    end = next((i for i in range(start, len(paragraphs)) if _SECTION.fullmatch(paragraphs[i].text.strip())), len(paragraphs))
    return [p for p in paragraphs[start:end] if p.text.strip()]


def sync_simple_fields(doc: DocxDocument, profile: CareerTruthProfile) -> list[SyncField]:
    fields: list[SyncField] = []
    all_paragraphs = _all_paragraphs(doc)
    contact = profile.contact_info
    for key, pattern in (("email", _EMAIL), ("phone", _PHONE)):
        desired = str(contact.get(key) or "").strip()
        if not desired:
            continue
        candidates = [p for p in all_paragraphs if pattern.search(p.text)]
        if len(candidates) != 1:
            _unplaced(fields, f"contact.{key}", desired, "contact line missing or ambiguous", blocking=len(candidates) > 1)
            continue
        p = candidates[0]
        wanted = pattern.sub(desired, p.text, count=1)
        _record_update(fields, f"contact.{key}", desired, p, wanted)

    # A uniquely identified name line near the top may be updated. If the old name cannot
    # be distinguished from an arbitrary heading, we report it instead of guessing.
    desired_name = str(contact.get("name") or "").strip()
    if desired_name:
        first = next((p for p in doc.paragraphs[:5] if p.text.strip()), None)
        if first is not None and not _SECTION.fullmatch(first.text.strip()) and not _EMAIL.search(first.text) and not _PHONE.search(first.text):
            _record_update(fields, "contact.name", desired_name, first, desired_name)
        else:
            _unplaced(fields, "contact.name", desired_name, "name line missing or ambiguous")

    desired_location = str(contact.get("location") or "").strip()
    if desired_location:
        contact_paragraphs = doc.paragraphs[:5]
        for section in doc.sections:
            contact_paragraphs.extend(section.header.paragraphs)
            contact_paragraphs.extend(section.footer.paragraphs)
        locations = [p for p in contact_paragraphs
                     if any(_LOCATION.fullmatch(part.strip()) for part in p.text.split("|"))
                     or re.match(r"^Location\s*:", p.text, re.I)]
        if len(locations) == 1:
            p = locations[0]
            prefix = re.match(r"^(Location\s*:\s*)", p.text, re.I)
            if prefix:
                wanted = prefix.group(1) + desired_location
            else:
                pieces = p.text.split("|")
                for index, piece in enumerate(pieces):
                    if _LOCATION.fullmatch(piece.strip()):
                        left = piece[:len(piece) - len(piece.lstrip())]
                        right = piece[len(piece.rstrip()):]
                        pieces[index] = left + desired_location + right
                        break
                wanted = "|".join(pieces)
            _record_update(fields, "contact.location", desired_location, p, wanted)
        else:
            _unplaced(fields, "contact.location", desired_location,
                      "location line missing or ambiguous", blocking=len(locations) > 1)

    if profile.summary.strip():
        summary = _section_body(doc, "SUMMARY") or _section_body(doc, "PROFESSIONAL SUMMARY")
        if not summary:
            structure = extract_docx_structure(doc)
            summary = [doc.paragraphs[i] for i in structure.summary_paragraph_indices]
        if not summary:
            first_heading = next((i for i, p in enumerate(doc.paragraphs)
                                  if _SECTION.fullmatch(p.text.strip())), len(doc.paragraphs))
            # Common header layout: name, contact, one summary paragraph, first section.
            if first_heading >= 3:
                summary = [p for p in doc.paragraphs[2:first_heading] if p.text.strip()]
        if len(summary) == 1:
            _record_update(fields, "summary", profile.summary, summary[0], profile.summary)
        else:
            _unplaced(fields, "summary", profile.summary, "summary paragraph missing or ambiguous")

    for key in ("skills", "tools"):
        values = getattr(profile, key)
        if not values:
            continue
        label = r"(?:Skills|Core Competencies)" if key == "skills" else "Tools"
        labelled = [p for p in doc.paragraphs if re.match(rf"^\s*{label}\s*:", p.text, re.I)]
        if len(labelled) == 1:
            p = labelled[0]
            prefix = re.match(rf"^(\s*{label}\s*:\s*)", p.text, re.I).group(1)
            _record_update(fields, key, ", ".join(values), p, prefix + ", ".join(values))
        else:
            headings = ("SKILLS", "TECHNICAL SKILLS", "CORE COMPETENCIES") if key == "skills" else ("TOOLS",)
            bodies = [_section_body(doc, heading) for heading in headings]
            targets = [body[0] for body in bodies if len(body) == 1 and ":" not in body[0].text]
            if len(targets) == 1:
                _record_update(fields, key, ", ".join(values), targets[0], ", ".join(values))
            else:
                _unplaced(fields, key, ", ".join(values), f"{key} line missing or ambiguous")
    return fields


def sync_education_and_certifications(doc: DocxDocument, profile: CareerTruthProfile) -> list[SyncField]:
    fields: list[SyncField] = []
    education = _section_body(doc, "EDUCATION")
    for entry in profile.education:
        candidates = [p for p in education if entry.institution.lower() in p.text.lower()]
        same_institution = [e for e in profile.education if e.institution.lower() == entry.institution.lower()]
        if len(candidates) > 1 and len(same_institution) > 1:
            prefix = f"{entry.degree} {entry.field}".strip().casefold()
            candidates = [p for p in candidates if p.text.split(",", 1)[0].strip().casefold() == prefix]
        if len(candidates) != 1:
            _unplaced(fields, "education", entry.institution, "education entry missing or ambiguous",
                      blocking=bool(education))
            continue
        p = candidates[0]
        # The resume parser can retain a whole education line as its institution
        # when it cannot split degree and school (decision 032). That line is
        # already represented in the file; do not try to parse it a second way.
        if not entry.degree and not entry.field and p.text.strip().casefold() == entry.institution.strip().casefold():
            if entry.year and str(entry.year) not in p.text:
                _unplaced(fields, "education", entry.institution, "education year is inconsistent", blocking=True)
            continue
        parts = [part.strip() for part in p.text.split(",")]
        if len(parts) < 3 or entry.institution.lower() not in parts[1].lower():
            _unplaced(fields, "education", entry.institution, "education format is ambiguous", blocking=True)
            continue
        wanted = f"{entry.degree} {entry.field}, {parts[1]}, {entry.year}"
        if len(parts) > 3:
            wanted += ", " + ", ".join(parts[3:])
        _record_update(fields, "education", entry.institution, p, wanted)

    if profile.certifications:
        body = _section_body(doc, "CERTIFICATIONS")
        if len(body) == 1:
            _record_update(fields, "certifications", ", ".join(profile.certifications), body[0],
                           ", ".join(profile.certifications))
        else:
            _unplaced(fields, "certifications", ", ".join(profile.certifications),
                      "certifications section missing or ambiguous", blocking=len(body) > 1)
    return fields
