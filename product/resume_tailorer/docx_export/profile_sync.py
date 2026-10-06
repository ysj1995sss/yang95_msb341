"""Bring a Word resume up to date with the Career Profile before tailoring (spec 011).

Tailoring edits the person's own Word file so its layout is kept (decision 006). The Career
Profile is where they fix and add facts, so without this step those edits never reached a
tailored resume. This updates a copy of the file to match the profile:

- roles are matched by employer and title, never by position;
- for a matched role, the file's bullet order is kept: an unchanged bullet is left byte-for-byte,
  an edited one is rewritten in place (same paragraph, same formatting), a bullet the profile no
  longer has is removed, and a new profile bullet is added after the role's last bullet with that
  bullet's style and numbering;
- a profile role the file doesn't have can't be placed in the original layout, and a file role
  the profile doesn't have is left as it is. Both are reported.

It also returns the profile reordered to the file's role order, because the bullet tailorer looks
roles up by their position in the file.
"""

from __future__ import annotations

import copy
import difflib
import io
import re
from dataclasses import dataclass, field, replace

from docx import Document

from resume_tailorer.models import CareerTruthProfile, WorkExperience
from resume_tailorer.parsers.bullets import strip_typed_bullet, typed_bullet_prefix
from resume_tailorer.parsers.docx_structure import extract_docx_structure

SAME_BULLET = 0.6  # similarity at which a profile bullet is taken as an edit of a file bullet


@dataclass
class SyncResult:
    docx_bytes: bytes
    aligned_profile: CareerTruthProfile  # work_experience in the file's role order
    changed: int = 0
    added: int = 0
    removed: int = 0
    unmatched_profile_roles: list[str] = field(default_factory=list)
    unmatched_file_roles: list[str] = field(default_factory=list)

    @property
    def summary(self) -> str:
        parts = []
        if self.changed or self.added or self.removed:
            parts.append(f"Brought your Career Profile into your Word file: {self.changed} bullet(s) updated, "
                         f"{self.added} added, {self.removed} removed.")
        if self.unmatched_profile_roles:
            parts.append("Not in your Word file, so not in this layout: " + "; ".join(self.unmatched_profile_roles)
                         + ". Add the role to your Word file, or tailor from a PDF to use a rebuilt layout.")
        if self.unmatched_file_roles:
            parts.append("In your Word file but not your Career Profile (left as it is): "
                         + "; ".join(self.unmatched_file_roles) + ".")
        return " ".join(parts)


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).strip()


def _role_label(job: WorkExperience) -> str:
    return " at ".join(p for p in (job.title, job.employer) if p) or "A role"


def _match_score(job: WorkExperience, block_text: str) -> float:
    text = _norm(block_text)
    employer, title = _norm(job.employer), _norm(job.title)
    score = 0.0
    if employer and employer in text:
        score += 2
    if title and title in text:
        score += 1
    return score


def _bullets(job: WorkExperience) -> list[str]:
    return [b.strip() for b in [*job.responsibilities, *job.accomplishments] if b and b.strip()]


def sync_docx_to_profile(docx_bytes: bytes, profile: CareerTruthProfile) -> SyncResult:
    doc = Document(io.BytesIO(docx_bytes))
    paragraphs = doc.paragraphs
    structure = extract_docx_structure(doc)

    # Match each file role to the profile role with the best employer/title match (one to one).
    block_texts = []
    for block in structure.jobs:
        parts = [paragraphs[block.anchor_index].text]
        if block.title_index is not None:
            parts.append(paragraphs[block.title_index].text)
        block_texts.append(" ".join(parts))
    pairs = sorted(((_match_score(job, text), b, j) for b, text in enumerate(block_texts)
                    for j, job in enumerate(profile.work_experience)), reverse=True)
    block_to_job: dict[int, int] = {}
    used_jobs: set[int] = set()
    for score, b, j in pairs:
        if score >= 2 and b not in block_to_job and j not in used_jobs:  # the employer must match
            block_to_job[b] = j
            used_jobs.add(j)

    result = SyncResult(docx_bytes=docx_bytes, aligned_profile=profile)
    result.unmatched_profile_roles = [_role_label(job) for j, job in enumerate(profile.work_experience)
                                      if j not in used_jobs]
    result.unmatched_file_roles = [" ".join(block_texts[b].split())[:80] for b in range(len(structure.jobs))
                                   if b not in block_to_job]

    # Work from the last role back, so inserting or removing paragraphs never shifts a role not yet done.
    for b in sorted(block_to_job, reverse=True):
        block = structure.jobs[b]
        job = profile.work_experience[block_to_job[b]]
        file_paras = [paragraphs[i] for i in block.bullet_paragraph_indices]
        wanted = _bullets(job)
        _sync_role(file_paras, wanted, result)

    # The profile in the file's role order, for the bullet tailorer (it looks roles up by their
    # position in the file). A file role the profile lacks keeps its place with its own text;
    # profile roles the file lacks go last.
    aligned_jobs = [profile.work_experience[block_to_job[b]] if b in block_to_job else
                    WorkExperience(employer="", title=block_texts[b], dates="", responsibilities=[], accomplishments=[])
                    for b in range(len(structure.jobs))]
    aligned_jobs += [job for j, job in enumerate(profile.work_experience) if j not in used_jobs]
    aligned = replace(profile, work_experience=aligned_jobs)
    result.aligned_profile = aligned

    if result.changed or result.added or result.removed:
        out = io.BytesIO()
        doc.save(out)
        result.docx_bytes = out.getvalue()
    return result


def _sync_role(file_paras: list, wanted: list[str], result: SyncResult) -> None:
    """Keep the file's bullet order; update, add and remove to match `wanted`."""
    current = [strip_typed_bullet(p.text).strip() for p in file_paras]
    unmatched_wanted = list(range(len(wanted)))
    assignment: dict[int, int] = {}  # file bullet index -> wanted index
    # Exact matches first, then the most similar remaining pairs.
    for f, text in enumerate(current):
        if text in wanted:
            w = next((i for i in unmatched_wanted if wanted[i] == text), None)
            if w is not None:
                assignment[f] = w
                unmatched_wanted.remove(w)
    candidates = sorted(((difflib.SequenceMatcher(None, current[f].lower(), wanted[w].lower()).ratio(), f, w)
                         for f in range(len(current)) if f not in assignment for w in unmatched_wanted),
                        reverse=True)
    for ratio, f, w in candidates:
        if ratio < SAME_BULLET:
            break
        if f in assignment or w not in unmatched_wanted:
            continue
        assignment[f] = w
        unmatched_wanted.remove(w)

    for f, paragraph in enumerate(file_paras):
        if f in assignment and current[f] != wanted[assignment[f]]:
            _set_text(paragraph, wanted[assignment[f]])
            result.changed += 1

    anchor = next((file_paras[f] for f in range(len(file_paras) - 1, -1, -1) if f in assignment), None)
    template = anchor or (file_paras[-1] if file_paras else None)
    if template is not None:
        last = anchor or template
        for w in sorted(unmatched_wanted):
            new_p = copy.deepcopy(template._p)
            last._p.addnext(new_p)
            from docx.text.paragraph import Paragraph

            paragraph = Paragraph(new_p, template._parent)
            _set_text(paragraph, wanted[w])
            last = paragraph
            result.added += 1
    for f, paragraph in enumerate(file_paras):
        if f not in assignment:
            paragraph._p.getparent().remove(paragraph._p)
            result.removed += 1


def _set_text(paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping its first run's formatting and any typed bullet."""
    prefix = typed_bullet_prefix(paragraph.text)
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(prefix + text)
        return
    runs[0].text = prefix + text
    for run in runs[1:]:
        run._element.getparent().remove(run._element)
