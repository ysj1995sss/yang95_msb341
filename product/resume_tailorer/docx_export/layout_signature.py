"""Structural signatures for proving DOCX template fidelity."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from typing import Iterable

from docx.document import Document as DocxDocument

from resume_tailorer.artifacts.models import (
    FindingCategory,
    FindingSeverity,
    ValidationFinding,
)


@dataclass(frozen=True)
class ParagraphLayoutSignature:
    identity: str
    style_id: str
    properties_xml: str


@dataclass(frozen=True)
class SectionLayoutSignature:
    properties_xml: str
    header_text_hash: str
    footer_text_hash: str


@dataclass(frozen=True)
class DocxLayoutSignature:
    sections: tuple[SectionLayoutSignature, ...]
    paragraphs: tuple[ParagraphLayoutSignature, ...]


def _text_hash(text: str) -> str:
    return sha256(text.encode("utf-8")).hexdigest()


def _part_text_hash(paragraphs: Iterable) -> str:
    return _text_hash("\n".join(paragraph.text for paragraph in paragraphs))


def capture_layout_signature(
    doc: DocxDocument,
    editable_paragraph_indices: set[int] | frozenset[int] = frozenset(),
) -> DocxLayoutSignature:
    """Capture immutable layout, masking text only for approved edit targets."""
    editable = set(editable_paragraph_indices)
    section_signatures = []
    for section in doc.sections:
        # Accessing header/footer lazily creates their relationship entries in
        # python-docx. Materialize them before reading sectPr so repeated
        # snapshots describe the same document state.
        header_text_hash = _part_text_hash(section.header.paragraphs)
        footer_text_hash = _part_text_hash(section.footer.paragraphs)
        section_signatures.append(
            SectionLayoutSignature(
                properties_xml=section._sectPr.xml,
                header_text_hash=header_text_hash,
                footer_text_hash=footer_text_hash,
            )
        )
    sections = tuple(section_signatures)
    paragraphs = tuple(
        ParagraphLayoutSignature(
            identity=(f"EDITABLE:{index}" if index in editable else _text_hash(paragraph.text)),
            style_id=(paragraph.style.style_id if paragraph.style is not None else ""),
            properties_xml=(paragraph._p.pPr.xml if paragraph._p.pPr is not None else ""),
        )
        for index, paragraph in enumerate(doc.paragraphs)
    )
    return DocxLayoutSignature(sections=sections, paragraphs=paragraphs)


def compare_layout_signatures(
    before: DocxLayoutSignature,
    after: DocxLayoutSignature,
) -> list[ValidationFinding]:
    findings: list[ValidationFinding] = []
    if before.sections != after.sections:
        findings.append(
            ValidationFinding(
                "DOCX_SECTION_LAYOUT_CHANGED",
                FindingSeverity.FAIL,
                FindingCategory.STRUCTURE,
                "Page, margin, header, footer, or section layout changed.",
            )
        )

    if len(before.paragraphs) != len(after.paragraphs):
        findings.append(
            ValidationFinding(
                "DOCX_PARAGRAPH_COUNT_CHANGED",
                FindingSeverity.FAIL,
                FindingCategory.STRUCTURE,
                "Paragraphs were added or removed from the original template.",
            )
        )
        return findings

    before_identity = tuple(item.identity for item in before.paragraphs)
    after_identity = tuple(item.identity for item in after.paragraphs)
    if before_identity != after_identity:
        findings.append(
            ValidationFinding(
                "DOCX_PARAGRAPH_ORDER_CHANGED",
                FindingSeverity.FAIL,
                FindingCategory.STRUCTURE,
                "A non-editable paragraph changed or paragraphs moved within the template.",
            )
        )

    before_format = tuple((item.style_id, item.properties_xml) for item in before.paragraphs)
    after_format = tuple((item.style_id, item.properties_xml) for item in after.paragraphs)
    if before_format != after_format:
        findings.append(
            ValidationFinding(
                "DOCX_PARAGRAPH_FORMAT_CHANGED",
                FindingSeverity.FAIL,
                FindingCategory.STRUCTURE,
                "Paragraph styles, numbering, indentation, tabs, spacing, or borders changed.",
            )
        )
    return findings
